"""inspect_excel — 기본은 읽기 전용 모드(셀 값·서식), --layout 일 때만 일반 모드(열 너비·행 높이·병합·페이지 설정)."""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "tools" / "office" / "inspect_excel.py"


def _load():
    spec = importlib.util.spec_from_file_location("inspect_excel_under_test", SCRIPT)
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture
def book(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "견적"
    ws["A1"] = "품명"
    ws["B2"] = 1234
    ws.merge_cells("A3:B3")
    ws.column_dimensions["A"].width = 20
    ws.row_dimensions[2].height = 30
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


def test_default_mode_prints_cells_but_not_layout(book, capsys):
    _load().inspect(str(book))
    out = capsys.readouterr().out
    assert "[ 셀 상세 ]" in out and "품명" in out and "1234" in out
    assert "[ 열 너비 ]" not in out and "[ 병합 셀 ]" not in out and "[ 페이지 설정 ]" not in out


def test_layout_mode_prints_widths_heights_merges_and_page_setup(book, capsys):
    _load().inspect(str(book), layout=True)
    out = capsys.readouterr().out
    assert "[ 열 너비 ]" in out and "A: width=20.00" in out
    assert "행002: height=30.0" in out
    assert "A3:B3" in out
    assert "[ 페이지 설정 ]" in out and "품명" in out


def test_default_mode_opens_the_workbook_read_only(book, monkeypatch):
    mod = _load()
    seen = {}
    real = mod.load_workbook

    def spy(path, **kwargs):
        seen.update(kwargs)
        return real(path, **kwargs)

    monkeypatch.setattr(mod, "load_workbook", spy)
    mod.inspect(str(book))
    assert seen.get("read_only") is True
    seen.clear()
    mod.inspect(str(book), layout=True)
    assert seen.get("read_only") is False
