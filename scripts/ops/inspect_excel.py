"""Excel 파일 셀 서식 전수 진단 스크립트.

사용법:
    python scripts/ops/inspect_excel.py <파일경로> [시트명]

출력:
    - 열 너비, 행 높이
    - 병합 셀 목록
    - 각 셀: 값, 폰트(크기/굵기), 정렬, 테두리, 배경색
    - 페이지 설정 (A4, 여백 등)
"""

import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


def inspect(file_path: str, sheet_name: str | None = None):
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    print(f"\n{'=' * 70}")
    print(f"파일: {file_path}")
    print(f"시트: {ws.title}  (max_row={ws.max_row}, max_col={ws.max_column})")
    print(f"{'=' * 70}")

    # ── 열 너비 ──────────────────────────────────────────────────────────
    print("\n[ 열 너비 ]")
    for col_letter, cd in sorted(ws.column_dimensions.items(), key=lambda x: column_index_from_string(x[0])):
        print(f"  {col_letter}: width={cd.width:.2f}  hidden={cd.hidden}")

    # ── 행 높이 ──────────────────────────────────────────────────────────
    print("\n[ 행 높이 ]")
    for r in range(1, ws.max_row + 1):
        rd = ws.row_dimensions[r]
        h = rd.height
        if h:
            print(f"  행{r:03d}: height={h:.1f}  hidden={rd.hidden}")

    # ── 병합 셀 ──────────────────────────────────────────────────────────
    print("\n[ 병합 셀 ]")
    for m in sorted(ws.merged_cells.ranges, key=lambda m: (m.min_row, m.min_col)):
        print(f"  {m.coord}  ({m.min_row}행{m.min_col}열 ~ {m.max_row}행{m.max_col}열)")

    # ── 셀 상세 ──────────────────────────────────────────────────────────
    print("\n[ 셀 상세 ]")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue

            addr = f"{get_column_letter(cell.column)}{cell.row}"

            # 폰트
            f = cell.font
            font_info = f"font(size={f.size}, bold={f.bold}, italic={f.italic}, color={f.color.rgb if f.color and f.color.type == 'rgb' else '-'})"

            # 정렬
            a = cell.alignment
            align_info = f"align(h={a.horizontal}, v={a.vertical}, wrap={a.wrap_text})"

            # 테두리
            b = cell.border

            def s(side):
                return side.border_style if side else None

            border_info = f"border(T={s(b.top)}, B={s(b.bottom)}, L={s(b.left)}, R={s(b.right)})"

            # 배경
            fill = cell.fill
            if fill and fill.fill_type == "solid":
                fg = fill.fgColor.rgb if fill.fgColor else "-"
                fill_info = f"fill({fg})"
            else:
                fill_info = "fill(-)"

            val = str(cell.value)[:30] if cell.value is not None else ""

            print(f"  {addr:6s} | {val:30s} | {font_info} | {align_info} | {border_info} | {fill_info}")

    # ── 페이지 설정 ──────────────────────────────────────────────────────
    print("\n[ 페이지 설정 ]")
    ps = ws.page_setup
    pm = ws.page_margins
    print(f"  paperSize={ps.paperSize} (9=A4)")
    print(f"  orientation={ps.orientation}")
    print(f"  fitToPage={getattr(ws.sheet_properties.pageSetUpPr, 'fitToPage', None)}")
    print(f"  fitToWidth={ps.fitToWidth}, fitToHeight={ps.fitToHeight}")
    print(f"  scale={ps.scale}")
    print(f"  margins: left={pm.left:.2f} right={pm.right:.2f} top={pm.top:.2f} bottom={pm.bottom:.2f}")
    print(f"  printArea={ws.print_area}")
    print(f"  showGridLines={ws.sheet_view.showGridLines}")
    print()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용법: python scripts/ops/inspect_excel.py <파일경로> [시트명]")
        sys.exit(1)
    inspect(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
