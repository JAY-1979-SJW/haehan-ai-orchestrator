"""건설공무 카페 키워드 필터 결과 → 기간별 시트 + 분석 Excel 출력."""

from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DATA_PATH = "data/cafe/keyword_filter_적산산출내역서단가.json"
OUT_PATH = "data/cafe/건설공무_적산산출내역서단가_분석.xlsx"
KEYWORDS = ["적산", "산출", "내역서", "단가"]

# ── 스타일 ────────────────────────────────────────────────────────────
H1_FILL = PatternFill("solid", fgColor="1F4E79")
H1_FONT = Font(color="FFFFFF", bold=True, size=10)
H2_FILL = PatternFill("solid", fgColor="2E75B6")
H2_FONT = Font(color="FFFFFF", bold=True, size=10)
H3_FILL = PatternFill("solid", fgColor="5B9BD5")
H3_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="BFBFBF")
BDR = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HLNK = Font(color="0563C1", underline="single", size=10)
BASE_FONT = Font(size=10)

KW_FILLS = {
    "적산": PatternFill("solid", fgColor="FFF2CC"),
    "산출": PatternFill("solid", fgColor="E2EFDA"),
    "내역서": PatternFill("solid", fgColor="DDEBF7"),
    "단가": PatternFill("solid", fgColor="FCE4D6"),
}
CAT_COLORS = {
    "계약·하도급": "D6E4F0",
    "공사관리": "D5F5E3",
    "기타": "F5F5F5",
    "안전": "FDEBD0",
    "장비·자재": "F9EBEA",
    "노무": "EAF2FF",
    "세무·회계": "FEF9E7",
    "법규·인허가": "F4ECF7",
    "행정·서류": "FDFEFE",
}
YELLOW = PatternFill("solid", fgColor="FFFACD")
GREEN = PatternFill("solid", fgColor="E8F5E9")
ORANGE = PatternFill("solid", fgColor="FFF3E0")


def _hdr(ws, row, cols, fill, font, heights=22):
    for ci, val in enumerate(cols, 1):
        c = ws.cell(row=row, column=ci, value=val)
        c.fill, c.font = fill, font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BDR
    ws.row_dimensions[row].height = heights


def _cell(ws, row, col, val, fill=None, font=None, align=None, bold=False):
    c = ws.cell(row=row, column=col, value=val)
    c.border = BDR
    c.font = Font(bold=bold, size=10) if not font else font
    c.alignment = align or Alignment(vertical="center")
    if fill:
        c.fill = fill
    return c


def _classify_type(title: str) -> str:
    q_words = ["?", "질문", "궁금", "어떻게", "어디서", "얼마", "어떤", "가능한가요", "되나요", "맞나요", "인가요"]
    ad_words = ["홍보", "모집", "강의", "학원", "교육", "무료", "안내드립니다", "알려드립니다", "합니다", "드립니다"]
    if any(w in title for w in q_words):
        return "질문"
    if any(w in title for w in ad_words):
        return "홍보/공지"
    return "정보/경험"


def _write_articles(ws, articles, start_row=2):
    cols = ["No", "날짜", "카테고리", "키워드", "유형", "제목", "조회수", "링크"]
    _hdr(ws, start_row, cols, H2_FILL, H2_FONT)

    for i, a in enumerate(articles, 1):
        r = start_row + i
        kw_str = "+".join(a.get("matched_kw", []))
        cat = a.get("category", "")
        title = a.get("title", "")
        atype = _classify_type(title)
        views = a.get("views", "") or ""
        href = a.get("href", "")
        bg = CAT_COLORS.get(cat)

        vals = [i, a.get("date", ""), cat, kw_str, atype, title, views, "링크"]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=val)
            c.border = BDR
            c.font = BASE_FONT
            c.alignment = Alignment(vertical="center", wrap_text=(ci == 6))
            if bg:
                c.fill = PatternFill("solid", fgColor=bg)

        kw_cell = ws.cell(row=r, column=4)
        for kw in KEYWORDS:
            if kw in kw_str:
                kw_cell.fill = KW_FILLS[kw]
                break

        if href:
            lc = ws.cell(row=r, column=8, value="링크")
            lc.hyperlink, lc.font = href, HLNK

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 52
    ws.column_dimensions["G"].width = 8
    ws.column_dimensions["H"].width = 6
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + len(articles) + 1


def _write_analysis(ws, articles, month_label, start_row):
    """월 분석 블록을 start_row 아래에 작성. 다음 빈 행 번호 반환."""
    r = start_row

    # ── 제목 ───────────────────────────────────────────────────────
    ws.merge_cells(f"A{r}:H{r}")
    c = ws.cell(row=r, column=1, value=f"📊  {month_label} 분석 요약")
    c.font = Font(bold=True, size=12, color="1F4E79")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r].height = 24
    r += 1

    # ── 키워드별 건수 ───────────────────────────────────────────────
    kw_c = Counter(kw for a in articles for kw in a.get("matched_kw", []))
    type_c = Counter(_classify_type(a.get("title", "")) for a in articles)

    _hdr(ws, r, ["키워드", "건수", "", "유형", "건수", "", "카테고리", "건수"], H3_FILL, H3_FONT)
    r += 1

    cat_c = Counter(a.get("category", "") for a in articles)
    kw_rows = [(kw, kw_c.get(kw, 0)) for kw in KEYWORDS]
    type_rows = list(type_c.most_common())
    cat_rows = list(cat_c.most_common())
    n_rows = max(len(kw_rows), len(type_rows), len(cat_rows))

    for i in range(n_rows):
        kw_v = kw_rows[i] if i < len(kw_rows) else ("", "")
        tp_v = type_rows[i] if i < len(type_rows) else ("", "")
        cat_v = cat_rows[i] if i < len(cat_rows) else ("", "")
        for ci, val in enumerate([kw_v[0], kw_v[1], "", tp_v[0], tp_v[1], "", cat_v[0], cat_v[1]], 1):
            cx = ws.cell(row=r, column=ci, value=val)
            cx.border = BDR
            cx.font = BASE_FONT
            cx.alignment = Alignment(horizontal="center" if ci in (2, 5, 8) else "left", vertical="center")
            # 키워드 색
            if ci == 1 and kw_v[0] in KW_FILLS:
                cx.fill = KW_FILLS[kw_v[0]]
            if ci == 4 and tp_v[0]:
                fx = YELLOW if tp_v[0] == "질문" else (GREEN if tp_v[0] == "정보/경험" else ORANGE)
                cx.fill = fx
            if ci == 7 and cat_v[0]:
                bg = CAT_COLORS.get(cat_v[0])
                if bg:
                    cx.fill = PatternFill("solid", fgColor=bg)
        r += 1

    r += 1

    # ── 조회수 TOP 10 ──────────────────────────────────────────────
    ws.merge_cells(f"A{r}:H{r}")
    ws.cell(row=r, column=1, value="🔥  조회수 TOP 10").font = Font(bold=True, size=11, color="C00000")
    ws.row_dimensions[r].height = 20
    r += 1

    _hdr(ws, r, ["순위", "날짜", "카테고리", "키워드", "유형", "제목", "조회수", "링크"], H3_FILL, H3_FONT)
    r += 1

    def _views_int(a):
        try:
            return int(str(a.get("views", 0) or 0).replace(",", ""))
        except Exception:
            return 0

    top10 = sorted(articles, key=_views_int, reverse=True)[:10]
    for rank, a in enumerate(top10, 1):
        kw_str = "+".join(a.get("matched_kw", []))
        cat = a.get("category", "")
        title = a.get("title", "")
        href = a.get("href", "")
        bg = CAT_COLORS.get(cat)

        vals = [rank, a.get("date", ""), cat, kw_str, _classify_type(title), title, _views_int(a), "링크"]
        for ci, val in enumerate(vals, 1):
            cx = ws.cell(row=r, column=ci, value=val)
            cx.border = BDR
            cx.font = BASE_FONT
            cx.alignment = Alignment(vertical="center", wrap_text=(ci == 6))
            if bg:
                cx.fill = PatternFill("solid", fgColor=bg)
        if href:
            lc = ws.cell(row=r, column=8, value="링크")
            lc.hyperlink, lc.font = href, HLNK
        r += 1

    r += 1

    # ── 주요 키워드 조합 패턴 ──────────────────────────────────────
    ws.merge_cells(f"A{r}:H{r}")
    ws.cell(row=r, column=1, value="📌  자주 나오는 키워드 조합").font = Font(bold=True, size=11, color="2E75B6")
    ws.row_dimensions[r].height = 20
    r += 1

    combo_c = Counter("+".join(sorted(a.get("matched_kw", []))) for a in articles)
    _hdr(ws, r, ["키워드 조합", "건수", "비율", "", "", "", "", ""], H3_FILL, H3_FONT)
    r += 1
    for combo, cnt in combo_c.most_common(8):
        ws.cell(row=r, column=1, value=combo).border = BDR
        ws.cell(row=r, column=2, value=cnt).border = BDR
        pct = ws.cell(row=r, column=3, value=f"{cnt / len(articles) * 100:.1f}%")
        pct.border = BDR
        pct.alignment = Alignment(horizontal="center")
        r += 1

    return r + 1


def main():
    data = json.loads(open(DATA_PATH, encoding="utf-8").read())

    for a in data:
        try:
            dt = datetime.strptime(a["date"], "%Y-%m-%d")
            a["_ym"] = dt.strftime("%Y-%m")
            a["_label"] = dt.strftime("%Y년 %m월")
            a["_quarter"] = f"{dt.year}년 {(dt.month - 1) // 3 + 1}분기"
        except Exception:
            a["_ym"] = "날짜미상"
            a["_label"] = "날짜미상"
            a["_quarter"] = "날짜미상"

    by_month: dict[str, list] = defaultdict(list)
    for a in data:
        by_month[a["_ym"]].append(a)

    months_sorted = sorted([m for m in by_month if m != "날짜미상"], reverse=True)
    if "날짜미상" in by_month:
        months_sorted.append("날짜미상")

    wb = Workbook()

    # ── 전체 요약 시트 ────────────────────────────────────────────────
    ws_s = wb.active
    ws_s.title = "전체요약"

    ws_s.merge_cells("A1:H1")
    c = ws_s["A1"]
    c.value = "건설공무 카페 — 적산·산출·내역서·단가 키워드 분석"
    c.font = Font(bold=True, size=15, color="1F4E79")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws_s.row_dimensions[1].height = 34

    ws_s.cell(
        row=2,
        column=1,
        value=f"수집 기간: {min(a['date'] for a in data)} ~ {max(a['date'] for a in data)}  |  총 {len(data)}건",
    ).font = Font(italic=True, color="595959", size=10)
    ws_s.merge_cells("A2:H2")
    ws_s["A2"].alignment = Alignment(horizontal="center")

    # 월별 요약
    _hdr(ws_s, 4, ["월", "건수", "내역서", "적산", "산출", "단가", "주요카테고리", "주요유형"], H1_FILL, H1_FONT)

    for ri, ym in enumerate(months_sorted, 5):
        arts = by_month[ym]
        kw_c = Counter(kw for a in arts for kw in a.get("matched_kw", []))
        cat_c = Counter(a.get("category", "") for a in arts)
        type_c = Counter(_classify_type(a.get("title", "")) for a in arts)
        top_c = cat_c.most_common(1)[0][0] if cat_c else ""
        top_t = type_c.most_common(1)[0][0] if type_c else ""
        label = arts[0]["_label"] if arts else ym
        row_vals = [
            label,
            len(arts),
            kw_c.get("내역서", 0),
            kw_c.get("적산", 0),
            kw_c.get("산출", 0),
            kw_c.get("단가", 0),
            top_c,
            top_t,
        ]
        for ci, val in enumerate(row_vals, 1):
            cx = ws_s.cell(row=ri, column=ci, value=val)
            cx.border = BDR
            cx.font = BASE_FONT
            cx.alignment = Alignment(horizontal="left" if ci in (1, 7, 8) else "center", vertical="center")

    total_row = 5 + len(months_sorted)
    kw_all = Counter(kw for a in data for kw in a.get("matched_kw", []))
    _hdr(
        ws_s,
        total_row,
        [
            "합계",
            len(data),
            kw_all.get("내역서", 0),
            kw_all.get("적산", 0),
            kw_all.get("산출", 0),
            kw_all.get("단가", 0),
            "",
            "",
        ],
        H1_FILL,
        H1_FONT,
    )

    ws_s.column_dimensions["A"].width = 14
    for col in "BCDEF":
        ws_s.column_dimensions[col].width = 10
    ws_s.column_dimensions["G"].width = 16
    ws_s.column_dimensions["H"].width = 12

    # 전체 TOP 15
    r_off = total_row + 2
    ws_s.merge_cells(f"A{r_off}:H{r_off}")
    ws_s.cell(row=r_off, column=1, value="🔥  전체 기간 조회수 TOP 15").font = Font(bold=True, size=12, color="C00000")
    ws_s.row_dimensions[r_off].height = 24
    r_off += 1
    _hdr(ws_s, r_off, ["순위", "날짜", "카테고리", "키워드", "유형", "제목", "조회수", "링크"], H1_FILL, H1_FONT)
    r_off += 1

    def _vi(a):
        try:
            return int(str(a.get("views", 0) or 0).replace(",", ""))
        except Exception:
            return 0

    for rank, a in enumerate(sorted(data, key=_vi, reverse=True)[:15], 1):
        kw_str = "+".join(a.get("matched_kw", []))
        cat = a.get("category", "")
        href = a.get("href", "")
        bg = CAT_COLORS.get(cat)
        title = a.get("title", "")
        vals = [rank, a.get("date", ""), cat, kw_str, _classify_type(title), title, _vi(a), "링크"]
        for ci, val in enumerate(vals, 1):
            cx = ws_s.cell(row=r_off, column=ci, value=val)
            cx.border = BDR
            cx.font = BASE_FONT
            cx.alignment = Alignment(vertical="center", wrap_text=(ci == 6))
            if bg:
                cx.fill = PatternFill("solid", fgColor=bg)
        if href:
            lc = ws_s.cell(row=r_off, column=8, value="링크")
            lc.hyperlink, lc.font = href, HLNK
        r_off += 1

    ws_s.freeze_panes = "A5"

    # ── 월별 시트 (목록 + 분석) ───────────────────────────────────────
    for ym in months_sorted:
        arts = by_month[ym]
        label = arts[0]["_label"] if arts else ym
        safe = label.replace(" ", "")[:20]
        ws = wb.create_sheet(title=safe)

        ws.merge_cells("A1:H1")
        c3 = ws["A1"]
        c3.value = f"{label}  ({len(arts)}건)"
        c3.font = Font(bold=True, size=13, color="1F4E79")
        c3.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # 분석 블록 먼저
        next_r = _write_analysis(ws, arts, label, start_row=2)

        # 구분선
        ws.merge_cells(f"A{next_r}:H{next_r}")
        ws.cell(row=next_r, column=1, value="▼  전체 게시글 목록").font = Font(bold=True, size=11, color="1F4E79")
        ws.row_dimensions[next_r].height = 20
        next_r += 1

        # 게시글 목록
        _write_articles(ws, sorted(arts, key=lambda x: x.get("date", ""), reverse=True), start_row=next_r)

    # ── 저장 ──────────────────────────────────────────────────────────
    wb.save(OUT_PATH)
    print(f"저장: {os.path.abspath(OUT_PATH)}")
    print(f"시트: {len(wb.sheetnames)}개 — {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
