"""Excel report generation for Naver SEO and shopping workflows."""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any


DATA_DIR = Path("data")
REPORT_DIR = DATA_DIR / "naver_reports"
LATEST_XLSX = REPORT_DIR / "naver_work_report_latest.xlsx"

DEFAULT_SOURCES = {
    "seo_plan": DATA_DIR / "naver_company_seo_plan_latest.json",
    "seo_diagnosis": DATA_DIR / "naver_company_seo_diagnosis_latest.json",
    "seo_assets": DATA_DIR / "naver_company_seo_assets_latest.json",
    "seo_ownership": DATA_DIR / "naver_company_seo_ownership_latest.json",
    "seo_exposure": DATA_DIR / "naver_company_seo_exposure_latest.json",
    "seo_submit_plan": DATA_DIR / "naver_company_seo_submit_plan_latest.json",
    "seo_monitor": DATA_DIR / "naver_company_seo_monitor_latest.json",
    "developers": DATA_DIR / "naver_developers_plan_latest.json",
    "shopping": DATA_DIR / "naver_shopping_competitors_latest.json",
    "catalog": DATA_DIR / "naver_service_action_catalog_latest.json",
}


def _load_json(path: Path) -> Any:
    if not path.exists():
        return {"missing": True, "path": str(path)}
    return json.loads(path.read_text(encoding="utf-8"))


def _safe_cell(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False)


def _short_sheet_name(name: str) -> str:
    return name[:31].replace("/", "_").replace("\\", "_").replace("?", "_")


def _write_mapping(ws, data: dict[str, Any], *, start_row: int = 1) -> int:
    row = start_row
    for key, value in data.items():
        if isinstance(value, dict):
            ws.cell(row=row, column=1, value=str(key))
            ws.cell(row=row, column=2, value=json.dumps(value, ensure_ascii=False))
        elif isinstance(value, list):
            ws.cell(row=row, column=1, value=str(key))
            ws.cell(row=row, column=2, value=json.dumps(value, ensure_ascii=False))
        else:
            ws.cell(row=row, column=1, value=str(key))
            ws.cell(row=row, column=2, value=_safe_cell(value))
        row += 1
    return row


def _write_list(ws, rows: list[Any], *, start_row: int = 1) -> int:
    if not rows:
        ws.cell(row=start_row, column=1, value="empty")
        return start_row + 1
    if all(isinstance(row, dict) for row in rows):
        keys: list[str] = []
        for row in rows:
            for key in row.keys():
                if key not in keys:
                    keys.append(key)
        for col, key in enumerate(keys, start=1):
            ws.cell(row=start_row, column=col, value=key)
        for ridx, item in enumerate(rows, start=start_row + 1):
            for col, key in enumerate(keys, start=1):
                ws.cell(row=ridx, column=col, value=_safe_cell(item.get(key)))
        return start_row + len(rows) + 1
    for ridx, item in enumerate(rows, start=start_row):
        ws.cell(row=ridx, column=1, value=_safe_cell(item))
    return start_row + len(rows)


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 80))
        ws.column_dimensions[letter].width = max(12, min(max_len + 2, 60))


def _write_source_sheet(wb, name: str, data: Any) -> None:
    ws = wb.create_sheet(_short_sheet_name(name))
    if isinstance(data, dict):
        if name == "shopping" and isinstance(data.get("items"), list):
            summary = {k: v for k, v in data.items() if k != "items"}
            row = _write_mapping(ws, summary)
            row += 1
            ws.cell(row=row, column=1, value="items")
            _write_list(ws, data["items"], start_row=row + 1)
        elif name == "seo_plan" and isinstance(data.get("phases"), list):
            summary = {k: v for k, v in data.items() if k != "phases"}
            row = _write_mapping(ws, summary)
            row += 1
            ws.cell(row=row, column=1, value="phases")
            _write_list(ws, data["phases"], start_row=row + 1)
        elif name == "seo_assets":
            summary = {k: v for k, v in data.items() if k not in ("diagnosis", "issues")}
            row = _write_mapping(ws, summary)
            row += 1
            ws.cell(row=row, column=1, value="issues")
            _write_list(ws, data.get("issues", []), start_row=row + 1)
        elif name in ("seo_ownership", "seo_submit_plan") and isinstance(data.get("methods") or data.get("submit_steps"), list):
            sequence_key = "methods" if "methods" in data else "submit_steps"
            summary = {k: v for k, v in data.items() if k != sequence_key}
            row = _write_mapping(ws, summary)
            row += 1
            ws.cell(row=row, column=1, value=sequence_key)
            _write_list(ws, data[sequence_key], start_row=row + 1)
        elif name == "seo_exposure" and isinstance(data.get("queries"), list):
            summary = {k: v for k, v in data.items() if k != "queries"}
            row = _write_mapping(ws, summary)
            row += 1
            ws.cell(row=row, column=1, value="queries")
            _write_list(ws, data["queries"], start_row=row + 1)
        elif name == "seo_monitor" and isinstance(data.get("checks"), list):
            summary = {k: v for k, v in data.items() if k != "checks"}
            row = _write_mapping(ws, summary)
            row += 1
            ws.cell(row=row, column=1, value="checks")
            _write_list(ws, data["checks"], start_row=row + 1)
        elif name == "catalog" and isinstance(data.get("features"), dict):
            _write_mapping(ws, {k: v for k, v in data.items() if k != "features"})
            rows = []
            for feature_name, feature in data["features"].items():
                rows.append({
                    "feature": feature_name,
                    "commands": ", ".join(feature.get("commands", [])),
                    "read": ", ".join(feature.get("read", [])),
                    "prepare": ", ".join(feature.get("prepare", [])),
                    "submit": ", ".join(feature.get("submit", [])),
                })
            _write_list(ws, rows, start_row=6)
        else:
            _write_mapping(ws, data)
    elif isinstance(data, list):
        _write_list(ws, data)
    else:
        ws.cell(row=1, column=1, value=_safe_cell(data))
    _autosize(ws)


def build_excel_report(
    *,
    output: str | Path | None = None,
    sources: dict[str, str | Path] | None = None,
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required to build Excel reports") from exc

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = Path(output) if output else REPORT_DIR / f"naver_work_report_{stamp}.xlsx"
    source_paths = {k: Path(v) for k, v in (sources or DEFAULT_SOURCES).items()}
    loaded = {name: _load_json(path) for name, path in source_paths.items()}

    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    summary_rows = [
        {"key": "generated_at", "value": datetime.now().isoformat(timespec="seconds")},
        {"key": "report_type", "value": "naver_seo_shopping_work_report"},
        {"key": "source_count", "value": len(source_paths)},
    ]
    for name, path in source_paths.items():
        summary_rows.append({
            "key": f"source.{name}",
            "value": str(path),
            "exists": path.exists(),
        })
    _write_list(summary, summary_rows)
    _autosize(summary)

    for name, data in loaded.items():
        _write_source_sheet(wb, name, data)

    wb.save(out_path)
    if out_path.resolve() != LATEST_XLSX.resolve():
        shutil.copyfile(out_path, LATEST_XLSX)

    return {
        "ok": True,
        "path": str(out_path),
        "latest_path": str(LATEST_XLSX),
        "sheet_count": len(wb.sheetnames),
        "sheets": wb.sheetnames,
        "sources": {name: str(path) for name, path in source_paths.items()},
    }


def print_excel_summary(result: dict[str, Any]) -> None:
    print("Naver Excel report")
    print(f"  path: {result['path']}")
    print(f"  latest: {result['latest_path']}")
    print(f"  sheets: {', '.join(result['sheets'])}")


__all__ = ["DEFAULT_SOURCES", "LATEST_XLSX", "build_excel_report", "print_excel_summary"]
