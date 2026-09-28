"""
scripts/hanafax/kras_campaign_template.py

"안전서류(KRAS) 전용 사이트 구축" 영업팩스 — 문서 생성.

사용자가 작성한 초안(2026-09-07)을 그대로 쓰되, "가나건설(주)"를
{{corp_nm}}으로, 두 지점(도입부·마무리)에 {{contract_name}}을 넣어
개인화한다. 표는 python-docx Table로 렌더링한다.

대상 데이터: g2b 프로젝트(C:\\work\\05. g2b) scripts/sales/scan_safety_doc_targets.py
가 만든 /nas/g2b/sales/scans/safety_doc_targets_*.json 을 로컬로 내려받아 사용한다.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

COMPARISON_ROWS = [
    ("기존 방식", "{corp_nm} 전용 AI 안전관리시스템", "기대효과"),
    ("현장마다 안전서류를 다시 찾고 작성", "126종 안전서류 작성·저장·출력", "반복작성 감소"),
    ("표준양식이 없는 업무는 직접 제작", "법적 요구사항을 반영한 자체 실무서식 제공", "양식 제작 부담 감소"),
    ("법령·고시를 필요할 때마다 검색", "AI로 관련 법령·조문·안전기준 질의", "확인시간 단축"),
    (
        "변경된 기준을 담당자가 직접 확인",
        "고용노동부·국가법령·KOSHA·국토부 기준 지속 확인 및 반영",
        "개정 대응 부담 감소",
    ),
    ("안전관리비 계획을 별도 작성·검토", "AI가 사용계획·사용항목·관련 근거 확인 지원", "계획업무 간소화"),
    ("현장별 자료가 PC·파일에 분산", "본사에서 현장별 안전자료 통합관리", "자료관리 편의"),
    ("현장 종료 후 자료가 흩어짐", "준공현장 자료 보관·조회", "회사 안전기록 축적"),
]

PRICE_ROWS = [
    ("귀사 전용 시스템 구축", "20만원 / 최초 1회"),
    ("본사 통합관리 및 유지관리", "월 5만원"),
    ("독립 운영현장", "현장당 월 3만원"),
]


def _h(doc, text, size=15, bold=True, color=(0x1D, 0x4E, 0x8C)):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(*color)
    return p


def _p(doc, text, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def _table(doc, rows, header_bold=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(val)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    if i == 0 and header_bold:
                        run.bold = True
    return t


def build_docx(corp_nm: str, contract_name: str, out_path: str) -> str:
    """업체별 개인화 팩스 docx 생성 → out_path 반환."""
    contract_name = contract_name or "이번 현장"
    doc = Document()

    _h(doc, f"{corp_nm}의 새로운 현장 준비에 도움이 되고자 연락드렸습니다", size=15)
    doc.add_paragraph()

    _p(
        doc,
        f"최근 귀사가 수주하신 「{contract_name}」 현장 관련 소식을 확인하던 중, "
        "새로운 현장을 준비하시는 시점에 도움이 될 수 있어 연락드렸습니다.",
    )
    _p(
        doc,
        "공사가 시작되면 위험성평가, TBM, 안전교육, 작업계획서, 각종 안전서류, "
        "산업안전보건관리비 사용계획, 관련 법령 확인 등 안전관리 업무도 함께 시작됩니다.",
    )
    doc.add_paragraph()

    _h(doc, f"해한AI가 {corp_nm}만을 위한 AI 안전관리시스템을 구축해드립니다.", size=13)
    _p(
        doc,
        "공용 사이트에 가입하는 방식이 아니라 귀사의 회사명·로고·현장정보를 적용한 "
        "전용 안전관리 환경으로 구성해드립니다.",
    )
    doc.add_paragraph()

    _h(doc, "기존 방식과 무엇이 달라집니까?", size=12)
    rows = [(a, b.format(corp_nm=corp_nm) if "{corp_nm}" in b else b, c) for a, b, c in COMPARISON_ROWS]
    _table(doc, rows)
    doc.add_paragraph()

    _h(doc, "안전기준은 계속 바뀝니다", size=12)
    _p(
        doc,
        "건설업 산업안전보건관리비 계상 및 사용기준만 보더라도 "
        "2023년 10월 → 2024년 9월 → 2025년 2월 연이어 개정되었습니다.",
    )
    _p(
        doc,
        "해한AI는 산업안전보건법령, 고용노동부 고시·질의회시, KOSHA 안전기준, "
        "국토교통부 건설안전 관련 기준을 지속적으로 확인하여 시스템에 반영해 나갑니다.",
    )
    _p(
        doc,
        "월 유지관리비는 단순한 서버비가 아닙니다 — 서버 운영 + 데이터 관리 + "
        "시스템 업데이트 + 법령·고시 모니터링 + 안전서류 유지관리 + AI 안전지식 업데이트가 포함됩니다.",
        size=10,
    )
    doc.add_paragraph()

    _h(doc, "비용", size=12)
    _table(doc, [("구분", "비용"), *list(PRICE_ROWS)])
    _p(doc, "본사에서 관리하는 소규모 현장은 하나의 회사 시스템에서 통합관리할 수 있습니다.", size=9.5)
    doc.add_paragraph()

    _p(
        doc,
        "월 5만원으로 단순히 안전서류를 사용하는 것이 아닙니다 — "
        "126종 안전서류 + 자체 실무서식 + AI 법령질의 + 안전관리비 관리 + 최신 기준 업데이트 + "
        f"회사·현장별 안전자료 관리를 하나의 {corp_nm} 전용 시스템으로 제공해드리는 것입니다.",
    )
    doc.add_paragraph()

    _h(doc, "한 현장을 위한 프로그램이 아니라 귀사의 안전관리 체계를 만들어드리는 것이 해한AI의 제안입니다.", size=12)
    _p(
        doc,
        f"「{contract_name}」 현장을 기준으로 먼저 보여드리겠습니다. "
        "직접 확인해보신 후 귀사에 도움이 될지 판단해주시면 됩니다.",
    )
    doc.add_paragraph()

    _h(doc, "해한AI · AI 안전관리시스템", size=12, color=(0x40, 0x40, 0x40))
    _p(doc, "안전관리시스템 : https://kras.haehan-ai.kr/")
    _p(doc, "전화 · 이메일 · 회사 홈페이지")
    doc.add_paragraph()
    note = doc.add_paragraph()
    r = note.add_run(
        "※ 본 시스템은 안전관리 업무의 작성·검색·분석 및 관련 기준 확인을 지원합니다. "
        "개별 현장의 법적 의무 및 산업안전보건관리비 사용 가능 여부는 실제 현장조건과 "
        "현행 기준에 따라 최종 확인이 필요합니다."
    )
    r.font.size = Pt(8)
    r.italic = True

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
