"""하나팩스 Playwright 기반 팩스 전송 엔진.

검증된 전송 흐름 (g2b 프로젝트 2026-03-31 실 전송 확인):
  1. 로그인 → Login 쿠키 확인
  2. 팩스보내기 메뉴 → tHanaFax_country.asp
  3. 수신번호 입력 + 추가
  4. 제목 입력
  5. docx 업로드 → TIF 변환 대기
  6. 팩스보내기 클릭
  7. submit_Result 확인 → 접수번호 추출
"""

from __future__ import annotations

import contextlib
import logging
import os
import re
import tempfile
import threading
from pathlib import Path
from typing import Any

from scripts.common.app_paths import repo_root

log = logging.getLogger("hanafax.sender")

ROOT = repo_root()
_TEMPLATE_DOCX = ROOT / "data" / "hanafax_template.docx"
_BASE_URL = "https://www.hanafax.com"
_LOCK = threading.Lock()

# 2026-07-30 고정: 영업팩스 공통 첨부파일 (무료수신거부 문구 반영본).
# 경로는 .env 의 HANAFAX_DEFAULT_ATTACH_FILE 로 관리 (PC마다 OneDrive 경로가 다름).
# 다른 파일로 바꾸려면 사용자가 명시적으로 새 경로를 지정할 것.
DEFAULT_ATTACH_FILE = os.environ.get("HANAFAX_DEFAULT_ATTACH_FILE", "")

# 팩스 첨부로 허용하는 확장자 — 임의 파일(설정/키 파일 등) 첨부 방지.
_ALLOWED_ATTACH_EXT = {".pdf", ".docx", ".doc"}


def _validate_attach_file(attach_file: str) -> str | None:
    """첨부파일 경로 검증. 문제 있으면 오류 메시지, 정상이면 None."""
    p = Path(attach_file).resolve()
    if not p.is_file():
        return f"첨부파일 없음: {attach_file}"
    if p.suffix.lower() not in _ALLOWED_ATTACH_EXT:
        return f"허용되지 않은 첨부파일 형식({p.suffix}): {attach_file}"
    return None


def send_fax(
    receiver_fax: str,
    subject: str,
    body: str,
    receiver_name: str = "",
    bid_name: str = "",
    user_id: str | None = None,
    password: str | None = None,
    attach_file: str | None = None,
) -> dict[str, Any]:
    """단건 팩스 전송.

    attach_file: 지정 시 docx 템플릿 대신 해당 파일(PDF 등)을 그대로 첨부.

    Returns:
        {"success": bool, "job_id": str|None, "message": str, "simulated": bool}
    """
    from scripts.hanafax.auth import get_credentials

    uid, pwd = user_id or "", password or ""
    if not uid or not pwd:
        uid, pwd = get_credentials()
    if not uid or not pwd:
        return {
            "success": False,
            "simulated": True,
            "job_id": None,
            "message": "자격증명 없음. python scripts/entry/cdp_cli.py cred set hanafax",
        }

    try:
        from playwright.sync_api import sync_playwright  # noqa
    except ImportError:
        return {
            "success": False,
            "simulated": False,
            "job_id": None,
            "message": "playwright 미설치. pip install playwright && playwright install chromium",
        }

    fax_no = re.sub(r"[^0-9]", "", receiver_fax)
    if len(fax_no) < 8:
        return {"success": False, "simulated": False, "job_id": None, "message": f"팩스번호 오류: {receiver_fax}"}

    if attach_file:
        err = _validate_attach_file(attach_file)
        if err:
            return {"success": False, "simulated": False, "job_id": None, "message": err}

    with _LOCK:
        return _run(uid, pwd, fax_no, subject, body, receiver_name, bid_name, attach_file)


def send_fax_bulk(
    receivers: list[dict],
    subject: str,
    attach_file: str | None = None,
    user_id: str | None = None,
    password: str | None = None,
) -> dict[str, Any]:
    """단체발송: 로그인·업로드 1회로 여러 수신번호에 동일 파일을 한 번에 전송.

    receivers: [{"receiver_fax": "02-...", "receiver_name": "..."}, ...]
    attach_file: 첨부할 실제 파일 경로 (PDF 등, 그대로 첨부 — docx 템플릿 미생성).
        생략 시 DEFAULT_ATTACH_FILE(영업팩스 공통 첨부) 사용.

    Returns:
        {"success": bool, "message": str, "sent_faxes": list[str], "job_id": str|None}
    """
    from scripts.hanafax.auth import get_credentials

    attach_file = attach_file or DEFAULT_ATTACH_FILE
    if not attach_file:
        return {
            "success": False,
            "sent_faxes": [],
            "job_id": None,
            "message": "첨부파일 미지정. attach_file 인자 또는 .env HANAFAX_DEFAULT_ATTACH_FILE 설정 필요",
        }
    uid, pwd = user_id or "", password or ""
    if not uid or not pwd:
        uid, pwd = get_credentials()
    if not uid or not pwd:
        return {
            "success": False,
            "sent_faxes": [],
            "job_id": None,
            "message": "자격증명 없음. python scripts/entry/cdp_cli.py cred set hanafax",
        }

    try:
        from playwright.sync_api import sync_playwright  # noqa
    except ImportError:
        return {
            "success": False,
            "sent_faxes": [],
            "job_id": None,
            "message": "playwright 미설치. pip install playwright && playwright install chromium",
        }

    err = _validate_attach_file(attach_file)
    if err:
        return {"success": False, "sent_faxes": [], "job_id": None, "message": err}

    fax_nos = []
    for r in receivers:
        fax_no = re.sub(r"[^0-9]", "", r.get("receiver_fax", ""))
        if len(fax_no) < 8:
            return {
                "success": False,
                "sent_faxes": [],
                "job_id": None,
                "message": f"팩스번호 오류: {r.get('receiver_fax')}",
            }
        fax_nos.append(fax_no)

    with _LOCK:
        return _run_bulk(uid, pwd, fax_nos, subject, attach_file)


def _run_bulk(uid: str, pwd: str, fax_nos: list[str], subject: str, attach_file: str) -> dict[str, Any]:
    from playwright.sync_api import TimeoutError as PWTimeout
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-setuid-sandbox", "--disable-dev-shm-usage"],
        )
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120",
            locale="ko-KR",
        )
        page = context.new_page()
        dialogs: list[str] = []
        page.on("dialog", lambda d: (dialogs.append(d.message), d.accept()))

        try:
            # STEP 1: 로그인
            page.goto(_BASE_URL, wait_until="domcontentloaded", timeout=20_000)
            page.fill('input[name="struid"]', uid)
            page.fill('input[name="strpwd"]', pwd)
            page.evaluate("document.querySelector('form').submit()")
            page.wait_for_load_state("domcontentloaded", timeout=15_000)
            page.wait_for_timeout(1_500)

            if not any(c["name"] == "Login" for c in context.cookies()):
                return {"success": False, "sent_faxes": [], "job_id": None, "message": f"로그인 실패. URL: {page.url}"}

            # STEP 2: 팩스 전송 페이지
            page.click('a[href*="tHanaFax_country"]')
            page.wait_for_load_state("domcontentloaded", timeout=15_000)
            page.wait_for_timeout(2_000)

            if "tHanaFax_country" not in page.url:
                return {
                    "success": False,
                    "sent_faxes": [],
                    "job_id": None,
                    "message": f"팩스 페이지 이동 실패: {page.url}",
                }

            # STEP 3: 수신번호 전체 추가 (단체발송)
            for fax_no in fax_nos:
                page.fill('input[name="inputfaxnumber"]', fax_no)
                page.click('button[onclick*="addPhoneNumber"]')
                page.wait_for_timeout(400)

            fax_list = page.evaluate("""
                () => Array.from(document.querySelector('select[name="faddList"]')?.options || [])
                    .map(o => o.value).filter(v => v !== 'none')
            """)
            added = set(fax_list)
            missing = [n for n in fax_nos if n not in added]
            if not fax_list:
                return {"success": False, "sent_faxes": [], "job_id": None, "message": "팩스번호 추가 실패 (0건 등록)"}

            # STEP 4: 제목
            page.fill('input[name="fax_title"]', subject[:80])

            # STEP 5: 파일 업로드 + TIF 변환 대기
            with page.expect_file_chooser(timeout=8_000) as fc_info:
                page.click('button[onclick*="pop_addfile"]')
            fc_info.value.set_files(attach_file)

            tif_file = None
            for _ in range(15):
                page.wait_for_timeout(2_000)
                tif_file = page.evaluate("""
                    () => {
                        const sel = document.getElementById('dropZone');
                        if (!sel || sel.options.length < 2) return null;
                        const val = sel.options[1].value;
                        return val && val.toLowerCase().endsWith('.tif') ? val : null;
                    }
                """)
                if tif_file:
                    break

            if not tif_file:
                tif_file = page.evaluate("""
                    () => { const sel = document.getElementById('dropZone');
                            return sel?.options[1]?.value || null; }
                """)
                if not tif_file:
                    return {"success": False, "sent_faxes": [], "job_id": None, "message": "TIF 변환 실패"}
                page.wait_for_timeout(5_000)

            # STEP 6: 팩스보내기 (변환중 다이얼로그 시 재시도)
            dialogs.clear()
            for retry in range(3):
                page.click('button[onclick*="e_money_chk"]')
                page.wait_for_timeout(3_000)
                if dialogs and any("변환" in d for d in dialogs):
                    dialogs.clear()
                    page.wait_for_timeout(5_000)
                    continue
                break

            if dialogs:
                err_kw = ["잔액", "부족", "오류", "실패", "충전"]
                if any(any(k in d for k in err_kw) for d in dialogs):
                    return {
                        "success": False,
                        "sent_faxes": [],
                        "job_id": None,
                        "message": f"전송 경고: {'; '.join(dialogs)}",
                    }

            # STEP 7: 결과 확인
            page.wait_for_load_state("domcontentloaded", timeout=30_000)
            page.wait_for_timeout(2_000)
            result_url = page.url

            if "submit_Result" not in result_url:
                body_txt = page.inner_text("body")[:200]
                return {
                    "success": False,
                    "sent_faxes": [],
                    "job_id": None,
                    "message": f"결과 페이지 미도달: {result_url} / {body_txt}",
                }

            result_text = page.inner_text("body")
            is_success = "팩스 전송 완료" in result_text

            job_id = None
            m = re.search(r"접수번호[\s:：]*([A-Z0-9\-]{5,})", result_text)
            if m:
                job_id = m.group(1)

            log.info(
                "단체발송 결과: success=%s job_id=%s 대상=%d건 미등록=%d건",
                is_success,
                job_id,
                len(fax_list),
                len(missing),
            )
            return {
                "success": is_success,
                "sent_faxes": fax_list,
                "missing_faxes": missing,
                "job_id": job_id,
                "message": "팩스 전송 완료" if is_success else f"결과 불명확: {result_text[:200]}",
            }

        except PWTimeout as e:
            return {"success": False, "sent_faxes": [], "job_id": None, "message": f"타임아웃: {e}"}
        except Exception as e:
            log.error("Playwright 오류(단체발송): %s", e, exc_info=True)
            return {"success": False, "sent_faxes": [], "job_id": None, "message": str(e)}
        finally:
            browser.close()


def _build_docx(subject: str, body: str, receiver_name: str, bid_name: str) -> str:
    """docx 템플릿 치환 후 임시 파일 경로 반환."""
    from datetime import datetime as _dt

    tmp_path = tempfile.mktemp(suffix=".docx", prefix="haehan_fax_")

    if not _TEMPLATE_DOCX.exists():
        log.warning("docx 템플릿 없음 — txt 폴백: %s", _TEMPLATE_DOCX)
        txt_path = tmp_path.replace(".docx", ".txt")
        with Path(txt_path).open("w", encoding="utf-8") as f:
            f.write(f"{subject}\n\n{body}")
        return txt_path

    try:
        from docx import Document

        doc = Document(str(_TEMPLATE_DOCX))
        now = _dt.now()

        _ORIG_BID = "율촌제1산단 정배수장 현대화사업 소방공사"
        _ORIG_BID_FULL = "율촌제1산단 정배수장 현대화사업 소방공사 소액수의 견적 제출 안내 공고"

        target_name = bid_name or subject
        replacements = {
            _ORIG_BID_FULL: target_name,
            _ORIG_BID: target_name.split(" 소액수의")[0] if " 소액수의" in target_name else target_name,
        }

        def _replace_runs(runs, old, new):
            full = "".join(r.text for r in runs)
            if old not in full:
                return
            runs[0].text = full.replace(old, new)
            for r in runs[1:]:
                r.text = ""

        for para in doc.paragraphs:
            for old, new in replacements.items():
                if old in para.text:
                    _replace_runs(para.runs, old, new)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for old, new in replacements.items():
                            if old in para.text:
                                _replace_runs(para.runs, old, new)

        if doc.tables:
            t0 = doc.tables[0]
            # 문서번호
            for para in t0.rows[0].cells[1].paragraphs:
                if "해한AI" in para.text:
                    new_no = f"해한AI {now.strftime('%y%m%d')}-{now.strftime('%H%M')}"
                    _replace_runs(para.runs, para.text.strip(), new_no)
            # 시행일자
            for para in t0.rows[0].cells[3].paragraphs:
                if para.text.strip():
                    _replace_runs(para.runs, para.text.strip(), now.strftime("%Y. %m. %d."))
            # 수신자
            if len(t0.rows) > 1:
                for para in t0.rows[1].cells[1].paragraphs:
                    txt = para.text.strip()
                    if txt:
                        _replace_runs(para.runs, txt, receiver_name or "수신자")
                    else:
                        para.add_run(receiver_name or "수신자")

        doc.save(tmp_path)
        log.info("docx 생성: receiver=%s bid=%s → %s", receiver_name, target_name, tmp_path)
        return tmp_path

    except ImportError:
        log.warning("python-docx 미설치 — txt 폴백")
        txt_path = tmp_path.replace(".docx", ".txt")
        with Path(txt_path).open("w", encoding="utf-8") as f:
            f.write(f"{subject}\n\n{body}")
        return txt_path


def _run(
    uid: str,
    pwd: str,
    fax_no: str,
    subject: str,
    body: str,
    receiver_name: str,
    bid_name: str,
    attach_file: str | None = None,
) -> dict[str, Any]:
    from playwright.sync_api import TimeoutError as PWTimeout
    from playwright.sync_api import sync_playwright

    using_external_file = bool(attach_file)
    tmp_path = attach_file if using_external_file else _build_docx(subject, body, receiver_name, bid_name)

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(
                headless=True,
                args=["--no-sandbox", "--disable-setuid-sandbox", "--disable-dev-shm-usage"],
            )
            context = browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120",
                locale="ko-KR",
            )
            page = context.new_page()
            dialogs: list[str] = []
            page.on("dialog", lambda d: (dialogs.append(d.message), d.accept()))

            try:
                # STEP 1: 로그인
                page.goto(_BASE_URL, wait_until="domcontentloaded", timeout=20_000)
                page.fill('input[name="struid"]', uid)
                page.fill('input[name="strpwd"]', pwd)
                page.evaluate("document.querySelector('form').submit()")
                page.wait_for_load_state("domcontentloaded", timeout=15_000)
                page.wait_for_timeout(1_500)

                if not any(c["name"] == "Login" for c in context.cookies()):
                    return {
                        "success": False,
                        "simulated": False,
                        "job_id": None,
                        "message": f"로그인 실패. URL: {page.url}",
                    }

                # STEP 2: 팩스 전송 페이지
                page.click('a[href*="tHanaFax_country"]')
                page.wait_for_load_state("domcontentloaded", timeout=15_000)
                page.wait_for_timeout(2_000)

                if "tHanaFax_country" not in page.url:
                    return {
                        "success": False,
                        "simulated": False,
                        "job_id": None,
                        "message": f"팩스 페이지 이동 실패: {page.url}",
                    }

                # STEP 3: 수신번호
                page.fill('input[name="inputfaxnumber"]', fax_no)
                page.click('button[onclick*="addPhoneNumber"]')
                page.wait_for_timeout(500)

                fax_list = page.evaluate("""
                    () => Array.from(document.querySelector('select[name="faddList"]')?.options || [])
                        .map(o => o.value).filter(v => v !== 'none')
                """)
                if not fax_list:
                    return {
                        "success": False,
                        "simulated": False,
                        "job_id": None,
                        "message": f"팩스번호 추가 실패: {fax_no}",
                    }

                # STEP 4: 제목
                page.fill('input[name="fax_title"]', subject[:80])

                # STEP 5: 파일 업로드 + TIF 변환 대기
                with page.expect_file_chooser(timeout=8_000) as fc_info:
                    page.click('button[onclick*="pop_addfile"]')
                fc_info.value.set_files(tmp_path)

                tif_file = None
                for _ in range(15):
                    page.wait_for_timeout(2_000)
                    tif_file = page.evaluate("""
                        () => {
                            const sel = document.getElementById('dropZone');
                            if (!sel || sel.options.length < 2) return null;
                            const val = sel.options[1].value;
                            return val && val.toLowerCase().endsWith('.tif') ? val : null;
                        }
                    """)
                    if tif_file:
                        break

                if not tif_file:
                    tif_file = page.evaluate("""
                        () => { const sel = document.getElementById('dropZone');
                                return sel?.options[1]?.value || null; }
                    """)
                    if not tif_file:
                        return {"success": False, "simulated": False, "job_id": None, "message": "TIF 변환 실패"}
                    page.wait_for_timeout(5_000)

                # STEP 6: 팩스보내기 (변환중 다이얼로그 시 재시도)
                dialogs.clear()
                for retry in range(3):
                    page.click('button[onclick*="e_money_chk"]')
                    page.wait_for_timeout(3_000)
                    if dialogs and any("변환" in d for d in dialogs):
                        dialogs.clear()
                        page.wait_for_timeout(5_000)
                        continue
                    break

                if dialogs:
                    err_kw = ["잔액", "부족", "오류", "실패", "충전"]
                    if any(any(k in d for k in err_kw) for d in dialogs):
                        return {
                            "success": False,
                            "simulated": False,
                            "job_id": None,
                            "message": f"전송 경고: {'; '.join(dialogs)}",
                        }

                # STEP 7: 결과 확인
                page.wait_for_load_state("domcontentloaded", timeout=20_000)
                page.wait_for_timeout(2_000)
                result_url = page.url

                if "submit_Result" not in result_url:
                    body_txt = page.inner_text("body")[:200]
                    return {
                        "success": False,
                        "simulated": False,
                        "job_id": None,
                        "message": f"결과 페이지 미도달: {result_url} / {body_txt}",
                    }

                result_text = page.inner_text("body")
                is_success = "팩스 전송 완료" in result_text

                job_id = None
                m = re.search(r"접수번호[\s:：]*([A-Z0-9\-]{5,})", result_text)
                if m:
                    job_id = m.group(1)
                else:
                    nums = re.findall(r"\b(\d{8,})\b", result_text)
                    if nums:
                        job_id = nums[0]

                log.info("전송 결과: success=%s job_id=%s fax=%s", is_success, job_id, fax_no)
                return {
                    "success": is_success,
                    "simulated": False,
                    "job_id": job_id,
                    "message": "팩스 전송 완료" if is_success else f"결과 불명확: {result_text[:200]}",
                }

            except PWTimeout as e:
                return {"success": False, "simulated": False, "job_id": None, "message": f"타임아웃: {e}"}
            except Exception as e:
                log.error("Playwright 오류: %s", e, exc_info=True)
                return {"success": False, "simulated": False, "job_id": None, "message": str(e)}
            finally:
                browser.close()
    finally:
        if not using_external_file:
            # 임시파일 정리 best-effort(실패해도 임시파일만 남을 뿐 안전, 팩스 발송 결과와 무관)
            with contextlib.suppress(Exception):
                Path(tmp_path).unlink()
