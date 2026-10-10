"""EUM 신규현장 설치대상 → Excel 출력."""

import json
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from scripts.common.app_paths import repo_root

ROOT = repo_root()
SRC = ROOT / "data" / "eum_new_sites_install_targets.json"
OUT_DIR = ROOT / "data"

COLUMNS = [
    ("NO", 8),
    ("공사번호", 16),
    ("공사명", 40),
    ("업체명", 22),
    ("현장주소", 30),
    ("관할지사", 14),
    ("담당자", 10),
    ("연락처", 16),
    ("이메일", 28),
    ("단말기설치 예정일", 16),
    ("단말기설치 예정대수", 14),
    ("공사시작일", 14),
    ("공사종료일", 14),
    ("등록일", 14),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
EVEN_FILL = PatternFill("solid", fgColor="EBF3FB")
ODD_FILL = PatternFill("solid", fgColor="FFFFFF")
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)


def make_excel(rows: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EUM 영업대상"
    ws.freeze_panes = "A2"

    # 헤더
    for col_idx, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 22

    # 데이터
    for row_idx, row in enumerate(rows, 2):
        fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL
        for col_idx, (col_name, _) in enumerate(COLUMNS, 1):
            val = row.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.border = BORDER
            cell.font = Font(size=9)
            # 공사명·주소·이메일은 좌측 정렬
            cell.alignment = LEFT if col_name in ("공사명", "현장주소", "이메일", "업체명") else CENTER
        ws.row_dimensions[row_idx].height = 16

    # 자동필터
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    # 요약 시트
    ws2 = wb.create_sheet("요약")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 12

    summary_data = [
        ("수집일시", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("전체 건수", len(rows)),
        ("이메일 보유", sum(1 for r in rows if r.get("이메일"))),
        ("연락처 보유", sum(1 for r in rows if r.get("연락처"))),
    ]

    # 관할지사별 집계
    from collections import Counter

    by_branch = Counter(r.get("관할지사", "미상") for r in rows)
    summary_data.append(("", ""))
    summary_data.append(("관할지사", "건수"))
    for branch, cnt in sorted(by_branch.items(), key=lambda x: -x[1]):
        summary_data.append((branch, cnt))

    for r_idx, (k, v) in enumerate(summary_data, 1):
        ws2.cell(row=r_idx, column=1, value=k).font = Font(
            bold=(k in ("수집일시", "전체 건수", "이메일 보유", "연락처 보유", "관할지사"))
        )
        ws2.cell(row=r_idx, column=2, value=v)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = OUT_DIR / f"eum_영업대상_{stamp}.xlsx"
    wb.save(out_path)
    return out_path


def main():
    rows = json.loads(SRC.read_text(encoding="utf-8"))
    print(f"로드: {len(rows)}건")
    out = make_excel(rows)
    size_kb = out.stat().st_size // 1024
    print(f"저장 완료: {out}")
    print(f"파일 크기: {size_kb} KB")


if __name__ == "__main__":
    main()
