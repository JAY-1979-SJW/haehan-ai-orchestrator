"""Excel 파일 셀 서식 전수 진단 스크립트.

사용법:
    python tools/office/inspect_excel.py <파일경로> [시트명]            # 셀 값·서식(읽기 전용 모드, 큰 파일도 가볍게)
    python tools/office/inspect_excel.py <파일경로> [시트명] --layout   # + 열 너비·행 높이·병합 셀·페이지 설정(일반 모드)

출력:
    - 각 셀: 값, 폰트(크기/굵기), 정렬, 테두리, 배경색
    - --layout 일 때만: 열 너비, 행 높이, 병합 셀 목록, 페이지 설정 (A4, 여백 등)
      (병합 셀·행 높이·페이지 설정은 openpyxl 읽기 전용 모드가 지원하지 않아 일반 모드로 연다)
"""

import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


def _print_layout(ws) -> None:
    """열 너비·행 높이·병합 셀 — 일반 모드에서만 읽을 수 있다."""
    print("\n[ 열 너비 ]")
    for col_letter, cd in sorted(ws.column_dimensions.items(), key=lambda x: column_index_from_string(x[0])):
        print(f"  {col_letter}: width={cd.width:.2f}  hidden={cd.hidden}")

    print("\n[ 행 높이 ]")
    for r in range(1, ws.max_row + 1):
        rd = ws.row_dimensions[r]
        h = rd.height
        if h:
            print(f"  행{r:03d}: height={h:.1f}  hidden={rd.hidden}")

    print("\n[ 병합 셀 ]")
    for m in sorted(ws.merged_cells.ranges, key=lambda m: (m.min_row, m.min_col)):
        print(f"  {m.coord}  ({m.min_row}행{m.min_col}열 ~ {m.max_row}행{m.max_col}열)")


def _side(side):
    return side.border_style if side else None


def _cell_line(cell) -> str:
    f = cell.font
    font_info = f"font(size={f.size}, bold={f.bold}, italic={f.italic}, color={f.color.rgb if f.color and f.color.type == 'rgb' else '-'})"
    a = cell.alignment
    align_info = f"align(h={a.horizontal}, v={a.vertical}, wrap={a.wrap_text})"
    b = cell.border
    border_info = f"border(T={_side(b.top)}, B={_side(b.bottom)}, L={_side(b.left)}, R={_side(b.right)})"
    fill = cell.fill
    if fill and fill.fill_type == "solid":
        fill_info = f"fill({fill.fgColor.rgb if fill.fgColor else '-'})"
    else:
        fill_info = "fill(-)"
    val = str(cell.value)[:30] if cell.value is not None else ""
    return f"  {cell.coordinate:6s} | {val:30s} | {font_info} | {align_info} | {border_info} | {fill_info}"


def _print_cells(ws) -> None:
    print("\n[ 셀 상세 ]")
    for row in ws.iter_rows():
        for cell in row:
            if not hasattr(cell, "coordinate"):  # 읽기 전용 모드의 빈 칸(EmptyCell)
                continue
            if cell.value is None and not getattr(cell, "has_style", False):
                continue
            print(_cell_line(cell))


def _print_page_setup(ws) -> None:
    print("\n[ 페이지 설정 ]")
    ps = ws.page_setup
    pm = ws.page_margins
    print(f"  paperSize={ps.paperSize} (9=A4)")
    print(f"  orientation={ps.orientation}")
    print(f"  fitToPage={getattr(ws.sheet_properties.pageSetUpPr, 'fitToPage', None)}")
    print(f"  fitToWidth={ps.fitToWidth}, fitToHeight={ps.fitToHeight}")
    print(f"  scale={ps.scale}")
    print(f"  margins: left={pm.left:.2f} right={pm.right:.2f} top={pm.top:.2f} bottom={pm.bottom:.2f}")
    print(f"  printArea={ws.print_area}")
    print(f"  showGridLines={ws.sheet_view.showGridLines}")
    print()


def inspect(file_path: str, sheet_name: str | None = None, *, layout: bool = False):
    """기본은 읽기 전용 모드로 셀 값·서식만, layout=True 면 일반 모드로 열어 배치(열 너비·행 높이·병합·페이지 설정)도 보여 준다."""
    wb = load_workbook(file_path, read_only=not layout)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        print(f"\n{'=' * 70}")
        print(f"파일: {file_path}")
        print(f"시트: {ws.title}  (max_row={ws.max_row}, max_col={ws.max_column})")
        print(f"{'=' * 70}")
        if layout:
            _print_layout(ws)
        _print_cells(ws)
        if layout:
            _print_page_setup(ws)
    finally:
        wb.close()


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--layout"]
    if not args:
        print("사용법: python tools/office/inspect_excel.py <파일경로> [시트명] [--layout]")
        sys.exit(1)
    inspect(args[0], args[1] if len(args) > 1 else None, layout="--layout" in sys.argv[1:])
