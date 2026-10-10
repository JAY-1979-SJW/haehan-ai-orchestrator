"""견적서 기준 템플릿 생성 — 서식만 있고 실제 데이터는 플레이스홀더.

이 파일을 엑셀에서 열어서:
  1. 열 너비 / 행 높이 원하는 대로 조정
  2. 테두리, 폰트, 배경색 조정
  3. 저장

그 다음:
  python tools/office/inspect_excel.py data/견적서_기준템플릿.xlsx --layout
로 서식 값을 추출하면 quote_generator.py 에 반영 가능.
"""

import sys

sys.path.insert(0, ".")

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path("data/견적서_기준템플릿.xlsx")

_THIN = Side(style="thin")
_MEDIUM = Side(style="medium")


def _all_thin():
    return Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)


def _center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def _left(wrap=False):
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)


def _right():
    return Alignment(horizontal="right", vertical="center")


def _gray():
    return PatternFill("solid", fgColor="D9E1F2")


def _lightgray():
    return PatternFill("solid", fgColor="F2F2F2")


wb = Workbook()
ws = wb.active
ws.title = "견적서"
ws.sheet_view.showGridLines = False

# ── A4 페이지 설정 ────────────────────────────────────────────────────────
ws.page_setup.paperSize = 9
ws.page_setup.orientation = "portrait"
ws.page_margins.left = 0.59
ws.page_margins.right = 0.59
ws.page_margins.top = 0.79
ws.page_margins.bottom = 0.79

# ══════════════════════════════════════════════════════════
# 열 너비 (초안 — 엑셀에서 조정 후 inspect 로 확정)
# A4 세로 기준: B~I 컬럼이 인쇄 영역
# ══════════════════════════════════════════════════════════
COL_W = {
    "A": 1.5,  # 좌 여백
    "B": 22,  # 품명
    "C": 7,  # 규격
    "D": 6,  # 수량
    "E": 13,  # 단가
    "F": 10,  # 기준(개월/구매)
    "G": 14,  # 공급가액
    "H": 12,  # 세액
    "I": 18,  # 비고
    "J": 1.5,  # 우 여백
}
for col, w in COL_W.items():
    ws.column_dimensions[col].width = w

# ══════════════════════════════════════════════════════════
# 행 높이 (초안)
# ══════════════════════════════════════════════════════════
ROW_H = {
    1: 45,  # 제목
    2: 8,  # 여백
    3: 22,  # 날짜
    4: 8,  # 여백
    5: 26,  # 수신처
    6: 22,  # 아래와 같이 견적합니다
    7: 8,  # 여백
    8: 22,  # 공급자: 등록번호
    9: 22,  # 공급자: 상호
    10: 22,  # 공급자: 대표
    11: 22,  # 공급자: 주소
    12: 22,  # 공급자: 업태
    13: 22,  # 공급자: 전화
    14: 8,  # 여백
    15: 34,  # 합계금액
    16: 8,  # 여백
    17: 28,  # 테이블 헤더
    18: 26,  # 품목1 (벽부형 임대)
    19: 26,  # 품목2 (벽부형 구매)
    20: 24,  # 빈행1
    21: 24,  # 빈행2
    22: 24,  # 빈행3
    23: 24,  # 빈행4
    24: 24,  # 빈행5
    25: 24,  # 빈행6
    26: 26,  # 소계
    27: 28,  # 합계
    28: 8,  # 여백
    29: 20,  # 안내문구1
    30: 20,  # 안내문구2
    31: 8,  # 하단 여백
}
for r, h in ROW_H.items():
    ws.row_dimensions[r].height = h


def mc(r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def C(r, c, val="", bold=False, size=10, align=None, border=None, fill=None, italic=False):  # noqa: PLR0913 - 공개 시그니처 유지(호출부 다수/CLI 인자 보존)
    cl = ws.cell(r, c, val)
    cl.font = Font(bold=bold, size=size, italic=italic)
    if align:
        cl.alignment = align
    if border:
        cl.border = border
    if fill:
        cl.fill = fill
    return cl


# ══════════════════════════════════════════════════════════
# ROW 1 — 제목
# ══════════════════════════════════════════════════════════
mc(1, 1, 1, 10)
C(1, 1, "견   적   서", bold=True, size=22, align=_center())

# ══════════════════════════════════════════════════════════
# ROW 3 — 날짜
# ══════════════════════════════════════════════════════════
mc(3, 2, 3, 5)
C(3, 2, "2026년  06월  23일", size=10, align=_left())

# ══════════════════════════════════════════════════════════
# ROW 5 — 수신처
# ══════════════════════════════════════════════════════════
mc(5, 2, 5, 5)
C(5, 2, "(주)아람정보통신  박원서  대표님", bold=True, size=13, align=_left())
mc(5, 6, 5, 9)
C(5, 6, "귀    하", bold=True, size=13, align=_center())

# ══════════════════════════════════════════════════════════
# ROW 6 — 견적 안내
# ══════════════════════════════════════════════════════════
mc(6, 2, 6, 9)
C(6, 2, "아래와 같이 견적합니다.", size=10, align=_left())

# ══════════════════════════════════════════════════════════
# ROW 8~13 — 공급자 블록
# ══════════════════════════════════════════════════════════
SUP_ROWS = [
    (8, "등 록 번 호", "372-34-00685"),
    (9, "상      호", "승민 F&G"),
    (10, "대      표", "신 재 우  (인)"),
    (11, "주      소", "서울시 강남구 봉은사로114길38 3층303호"),
    (12, "업      태", "건설업, 도매 및 소매  /  종목: 전기통신, 소방자재"),
    (13, "전      화", "02-562-6652  /  팩스: 02-6442-6665"),
]
for r, lbl, val in SUP_ROWS:
    C(r, 6, lbl, bold=True, size=10, align=_center(), border=_all_thin(), fill=_lightgray())
    mc(r, 7, r, 9)
    C(r, 7, val, size=10, align=_left(wrap=True), border=_all_thin())
    # 병합 범위 경계 셀 테두리 보정
    ws.cell(r, 9).border = Border(top=_THIN, bottom=_THIN, right=_THIN)

# ══════════════════════════════════════════════════════════
# ROW 15 — 합계금액 박스
# ══════════════════════════════════════════════════════════
mc(15, 2, 15, 5)
C(
    15,
    2,
    "합  계  금  액\n(공급가액 + 세액)",
    bold=True,
    size=11,
    align=_center(wrap=True),
    border=_all_thin(),
    fill=_lightgray(),
)
mc(15, 6, 15, 9)
C(
    15,
    6,
    "₩  3,168,000  원",
    bold=True,
    size=14,
    align=_center(),
    border=Border(top=_MEDIUM, bottom=_MEDIUM, left=_MEDIUM, right=_MEDIUM),
)
# 경계 보정
ws.cell(15, 5).border = Border(top=_THIN, bottom=_THIN, right=_THIN)
ws.cell(15, 9).border = Border(top=_MEDIUM, bottom=_MEDIUM, right=_MEDIUM)

# ══════════════════════════════════════════════════════════
# ROW 17 — 테이블 헤더
# ══════════════════════════════════════════════════════════
HDR = [
    ("B", "품      명"),
    ("C", "규격"),
    ("D", "수량"),
    ("E", "단    가"),
    ("F", "기    준"),
    ("G", "공 급 가 액"),
    ("H", "세    액"),
    ("I", "비    고"),
]
for col_l, txt in HDR:
    c = ord(col_l) - ord("A") + 1
    C(17, c, txt, bold=True, size=10, align=_center(wrap=True), border=_all_thin(), fill=_gray())

# ══════════════════════════════════════════════════════════
# ROW 18~19 — 품목 샘플
# ══════════════════════════════════════════════════════════
ITEMS = [
    (18, "벽부형(ACR-C) 본체", "EA", 1, "70,000", "24개월", "1,680,000", "168,000", ""),
    (19, "벽부형(ACR-C) 본체", "EA", 1, "1,200,000", "구  매", "1,200,000", "120,000", ""),
]
for r, nm, spec, qty, up, kijun, supply, tax, bigo in ITEMS:
    vals = [nm, spec, qty, up, kijun, supply, tax, bigo]
    for i, (col_l, _) in enumerate(HDR):
        c = ord(col_l) - ord("A") + 1
        align = _left(wrap=True) if col_l == "I" else _center()
        C(r, c, vals[i], size=10, align=align, border=_all_thin())

# ══════════════════════════════════════════════════════════
# ROW 20~25 — 빈 품목 행 (서식만)
# ══════════════════════════════════════════════════════════
for r in range(20, 26):
    for col_l, _ in HDR:
        c = ord(col_l) - ord("A") + 1
        ws.cell(r, c).border = _all_thin()

# ══════════════════════════════════════════════════════════
# ROW 26 — 소계
# ══════════════════════════════════════════════════════════
mc(26, 2, 26, 6)
C(26, 2, "소          계", bold=True, size=10, align=_center(), border=_all_thin(), fill=_lightgray())
# 병합 경계 보정
for c in range(3, 7):
    ws.cell(26, c).border = Border(top=_THIN, bottom=_THIN)
ws.cell(26, 6).border = Border(top=_THIN, bottom=_THIN, right=_THIN)
C(26, 7, "2,880,000", size=10, align=_right(), border=_all_thin())
C(26, 8, "288,000", size=10, align=_right(), border=_all_thin())
C(26, 9, "", border=_all_thin())

# ══════════════════════════════════════════════════════════
# ROW 27 — 합계
# ══════════════════════════════════════════════════════════
mc(27, 2, 27, 6)
C(27, 2, "합          계", bold=True, size=11, align=_center(), border=_all_thin(), fill=_gray())
for c in range(3, 7):
    ws.cell(27, c).border = Border(top=_THIN, bottom=_THIN)
ws.cell(27, 6).border = Border(top=_THIN, bottom=_THIN, right=_THIN)
mc(27, 7, 27, 9)
C(27, 7, "₩  3,168,000  원", bold=True, size=11, align=_center(), border=_all_thin())
ws.cell(27, 9).border = Border(top=_THIN, bottom=_THIN, right=_THIN)

# ══════════════════════════════════════════════════════════
# ROW 29~30 — 안내 문구
# ══════════════════════════════════════════════════════════
mc(29, 2, 29, 9)
C(
    29,
    2,
    "※ 임대형은 무상 1년 A/S 적용, 구매형은 무상 1년 A/S 후 100% 환급 적용됩니다.",
    size=9,
    italic=True,
    align=_left(),
)
mc(30, 2, 30, 9)
C(30, 2, "※ 위 견적금액은 부가세(VAT 10%) 포함 금액입니다.", size=9, italic=True, align=_left())

# ── 인쇄 영역 ────────────────────────────────────────────────────────────
ws.print_area = "A1:J31"

wb.save(OUT)
print(f"템플릿 생성: {OUT.resolve()}")
print()
print("다음 단계:")
print("  1. 파일을 엑셀에서 열어 열너비/행높이/테두리/폰트 등 원하는 대로 조정")
print("  2. 저장")
print("  3. python tools/office/inspect_excel.py data/견적서_기준템플릿.xlsx --layout")
print("  4. 출력된 값을 quote_generator.py 에 반영")
