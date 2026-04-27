"""Excel 커넥터 1단계 테스트.

검증:
1) excel_read_sheet 기본 성공
2) 특정 sheet 읽기 성공
3) 없는 sheet 차단 → SHEET_NOT_FOUND
4) 허용되지 않은 경로 차단 → FILE_NOT_ALLOWED
5) 경로 탈출 차단 (work 디렉터리 바깥)
6) 원본 overwrite 차단 (output_path == source_path)
7) output_path 허용 디렉터리 외 차단
8) 공통 response 형식 ({ok, action, data, error})
9) 표준 error code 검증
10) 로그 공통 필드 (action, category=excel, risk_level, ok, duration_ms, timestamp)
"""
from __future__ import annotations

import json
import os
import sys

import pytest

sys.path.insert(
    0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
)


_COMMON_KEYS = {"ok", "action", "data", "error"}

_LOG_COMMON_KEYS = {
    "action", "category", "risk_level",
    "site_key", "target_url",
    "ok", "duration_ms", "timestamp",
}


@pytest.fixture()
def env(tmp_path, monkeypatch):
    """격리된 work / output / log 디렉터리를 설정한다."""
    from agent import config as _cfg
    from agent import runner as _runner

    work_dir = tmp_path / "work"
    out_dir = tmp_path / "output" / "reports"
    log_path = tmp_path / "agent_actions.jsonl"
    work_dir.mkdir(parents=True)
    out_dir.mkdir(parents=True)

    monkeypatch.setattr(_cfg, "AGENT_WORK_DIR", work_dir)
    monkeypatch.setattr(_cfg, "AGENT_OUTPUT_DIR", out_dir)
    monkeypatch.setattr(_cfg, "AGENT_LOG_PATH", log_path)

    if _runner._browser_lock.locked():
        try:
            _runner._browser_lock.release()
        except RuntimeError:
            pass

    yield {
        "tmp": tmp_path,
        "work": work_dir,
        "out": out_dir,
        "log": log_path,
    }

    if _runner._browser_lock.locked():
        try:
            _runner._browser_lock.release()
        except RuntimeError:
            pass


def _read_log(log_path):
    if not log_path.exists():
        return []
    return [
        json.loads(ln)
        for ln in log_path.read_text(encoding="utf-8").strip().splitlines()
        if ln
    ]


def _make_xlsx(path, sheets: dict[str, list[list]]):
    """간단한 xlsx 를 생성. sheets 는 {sheet_name: rows}."""
    from openpyxl import Workbook
    wb = Workbook()
    default = wb.active
    first_name, first_rows = next(iter(sheets.items()))
    default.title = first_name
    for r in first_rows:
        default.append(r)
    for name, rows in list(sheets.items())[1:]:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


# ══════════════════════════════════════════════════════════════════════
# 1) excel_read_sheet 기본 성공
# ══════════════════════════════════════════════════════════════════════
def test_excel_read_sheet_success(env):
    from agent import app

    src = env["work"] / "a.xlsx"
    _make_xlsx(src, {"Sheet1": [["a", "b"], [1, 2], [3, 4]]})

    out = app.run(
        "excel_read_sheet",
        options={"file_path": str(src)},
    )
    assert set(out.keys()) == _COMMON_KEYS
    assert out["ok"] is True
    assert out["action"] == "excel_read_sheet"
    assert out["error"] is None
    data = out["data"]
    assert data["sheet_name"] == "Sheet1"
    assert data["row_count"] == 3
    assert data["column_count"] == 2
    assert data["rows"][0] == ["a", "b"]
    assert data["rows"][1] == [1, 2]


# ══════════════════════════════════════════════════════════════════════
# 2) 특정 sheet 읽기 성공
# ══════════════════════════════════════════════════════════════════════
def test_excel_read_sheet_specific_sheet(env):
    from agent import app

    src = env["work"] / "b.xlsx"
    _make_xlsx(src, {
        "First": [["x"]],
        "Target": [["hello", "world"], [10, 20]],
    })

    out = app.run(
        "excel_read_sheet",
        options={"file_path": str(src), "sheet_name": "Target"},
    )
    assert out["ok"] is True
    assert out["data"]["sheet_name"] == "Target"
    assert out["data"]["rows"][0] == ["hello", "world"]
    assert out["data"]["row_count"] == 2


# ══════════════════════════════════════════════════════════════════════
# 3) 없는 sheet 차단
# ══════════════════════════════════════════════════════════════════════
def test_excel_read_sheet_missing_sheet_blocked(env):
    from agent import app, errors

    src = env["work"] / "c.xlsx"
    _make_xlsx(src, {"Only": [["1"]]})

    out = app.run(
        "excel_read_sheet",
        options={"file_path": str(src), "sheet_name": "DoesNotExist"},
    )
    assert out["ok"] is False
    assert out["error"] == errors.SHEET_NOT_FOUND
    assert errors.is_standard_code(out["error"])


# ══════════════════════════════════════════════════════════════════════
# 4) 허용되지 않은 경로 차단 (work 디렉터리 밖 절대경로)
# ══════════════════════════════════════════════════════════════════════
def test_excel_read_sheet_outside_work_dir_blocked(env):
    from agent import app, errors

    outside = env["tmp"] / "outside.xlsx"
    _make_xlsx(outside, {"S": [["x"]]})

    out = app.run(
        "excel_read_sheet",
        options={"file_path": str(outside)},
    )
    assert out["ok"] is False
    assert out["error"] == errors.FILE_NOT_ALLOWED
    assert errors.is_standard_code(out["error"])


# ══════════════════════════════════════════════════════════════════════
# 5) 경로탈출 (..) 차단
# ══════════════════════════════════════════════════════════════════════
def test_excel_read_sheet_path_traversal_blocked(env):
    from agent import app, errors

    sibling = env["tmp"] / "sibling.xlsx"
    _make_xlsx(sibling, {"S": [["x"]]})

    # work 디렉터리에서 .. 를 이용해 sibling 접근
    traversal = str(env["work"] / ".." / "sibling.xlsx")
    out = app.run(
        "excel_read_sheet",
        options={"file_path": traversal},
    )
    assert out["ok"] is False
    assert out["error"] == errors.FILE_NOT_ALLOWED


def test_excel_read_sheet_relative_path_blocked(env):
    from agent import app, errors

    out = app.run(
        "excel_read_sheet",
        options={"file_path": "a.xlsx"},
    )
    assert out["ok"] is False
    assert out["error"] == errors.FILE_NOT_ALLOWED


def test_excel_read_sheet_empty_path_blocked(env):
    from agent import app, errors

    out = app.run("excel_read_sheet", options={})
    assert out["ok"] is False
    assert out["error"] == errors.FILE_PATH_REQUIRED


def test_excel_read_sheet_missing_file(env):
    from agent import app, errors

    ghost = env["work"] / "ghost.xlsx"
    out = app.run("excel_read_sheet", options={"file_path": str(ghost)})
    assert out["ok"] is False
    assert out["error"] == errors.FILE_NOT_FOUND


# ══════════════════════════════════════════════════════════════════════
# 6) 원본 overwrite 차단
# ══════════════════════════════════════════════════════════════════════
def test_excel_write_source_overwrite_blocked(env):
    from agent import app, errors

    src = env["work"] / "src.xlsx"
    _make_xlsx(src, {"S": [["x"]]})

    out = app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            # 출력 경로가 원본과 동일 → 차단되어야 한다.
            "output_path": str(src),
            "rows": [["h1"], [1]],
        },
    )
    assert out["ok"] is False
    assert out["error"] == errors.OUTPUT_PATH_NOT_ALLOWED


def test_excel_write_existing_output_not_overwritten(env):
    from agent import app, errors

    src = env["work"] / "src.xlsx"
    _make_xlsx(src, {"S": [["x"]]})

    existing = env["out"] / "exists.xlsx"
    existing.write_bytes(b"already here")

    out = app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": str(existing),
            "rows": [["h"]],
        },
    )
    assert out["ok"] is False
    assert out["error"] == errors.OUTPUT_FILE_EXISTS
    # 기존 내용 보존 확인
    assert existing.read_bytes() == b"already here"


# ══════════════════════════════════════════════════════════════════════
# 7) output_path 허용 디렉터리 외 차단
# ══════════════════════════════════════════════════════════════════════
def test_excel_write_output_outside_allowed_dir(env):
    from agent import app, errors

    src = env["work"] / "src.xlsx"
    _make_xlsx(src, {"S": [["x"]]})

    outside_out = env["tmp"] / "outside_out.xlsx"
    out = app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": str(outside_out),
            "rows": [["h"]],
        },
    )
    assert out["ok"] is False
    assert out["error"] == errors.OUTPUT_PATH_NOT_ALLOWED
    assert not outside_out.exists()


def test_excel_write_output_traversal_blocked(env):
    from agent import app, errors

    src = env["work"] / "src.xlsx"
    _make_xlsx(src, {"S": [["x"]]})

    traversal = str(env["out"] / ".." / "escape.xlsx")
    out = app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": traversal,
            "rows": [["h"]],
        },
    )
    assert out["ok"] is False
    assert out["error"] == errors.OUTPUT_PATH_NOT_ALLOWED


def test_excel_write_report_copy_success(env):
    from agent import app, errors

    src = env["work"] / "src.xlsx"
    _make_xlsx(src, {"S": [["x"]]})

    target = env["out"] / "report.xlsx"
    rows = [["col1", "col2"], [1, 2], [3, 4]]

    out = app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": str(target),
            "rows": rows,
        },
    )
    assert out["ok"] is True
    assert out["error"] is None
    data = out["data"]
    assert data["written_row_count"] == 3
    assert data["output_path"].endswith("report.xlsx")
    assert target.exists()

    # 원본은 바뀌지 않았어야 한다
    from openpyxl import load_workbook
    wb = load_workbook(str(src), read_only=True)
    try:
        assert wb.sheetnames == ["S"]
        assert list(wb["S"].iter_rows(values_only=True)) == [("x",)]
    finally:
        wb.close()


# ══════════════════════════════════════════════════════════════════════
# 8) 공통 response 형식
# ══════════════════════════════════════════════════════════════════════
def test_common_response_shape_for_read_and_write(env):
    from agent import app

    src = env["work"] / "s.xlsx"
    _make_xlsx(src, {"S": [["a"]]})

    r1 = app.run("excel_read_sheet", options={"file_path": str(src)})
    assert set(r1.keys()) == _COMMON_KEYS
    assert r1["action"] == "excel_read_sheet"
    assert isinstance(r1["data"], dict)

    r2 = app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": str(env["out"] / "ok.xlsx"),
            "rows": [["x"]],
        },
    )
    assert set(r2.keys()) == _COMMON_KEYS
    assert r2["action"] == "excel_write_report_copy"


# ══════════════════════════════════════════════════════════════════════
# 9) 표준 error code 검증
# ══════════════════════════════════════════════════════════════════════
def test_excel_standard_error_codes_exist():
    from agent import errors

    for name in (
        "FILE_PATH_REQUIRED",
        "FILE_NOT_ALLOWED",
        "FILE_NOT_FOUND",
        "SHEET_NOT_FOUND",
        "OUTPUT_PATH_NOT_ALLOWED",
        "OUTPUT_FILE_EXISTS",
        "EXCEL_UNSUPPORTED_FORMAT",
        "ROWS_REQUIRED",
    ):
        assert hasattr(errors, name), f"errors.{name} missing"
        assert getattr(errors, name) in errors.ALL_CODES, name

    assert errors.EXCEL_ERROR in errors.PREFIX_CODES
    assert errors.is_standard_code("excel_error:IOError")
    assert errors.is_standard_code(errors.SHEET_NOT_FOUND)


def test_excel_unsupported_format_blocked(env):
    from agent import app, errors

    src = env["work"] / "bad.xlsm"
    src.write_bytes(b"not really xlsm")
    out = app.run("excel_read_sheet", options={"file_path": str(src)})
    assert out["ok"] is False
    assert out["error"] == errors.EXCEL_UNSUPPORTED_FORMAT
    assert errors.is_standard_code(out["error"])


# ══════════════════════════════════════════════════════════════════════
# 10) 로그 공통 필드
# ══════════════════════════════════════════════════════════════════════
def test_log_common_fields_for_excel_read(env):
    from agent import app

    src = env["work"] / "L.xlsx"
    _make_xlsx(src, {"S": [["1"]]})
    app.run("excel_read_sheet", options={"file_path": str(src)})

    entries = _read_log(env["log"])
    assert entries, "no log entry emitted"
    last = entries[-1]
    assert _LOG_COMMON_KEYS.issubset(last.keys()), (
        f"missing keys: {_LOG_COMMON_KEYS - set(last.keys())}"
    )
    assert last["action"] == "excel_read_sheet"
    assert last["category"] == "excel"
    assert last["risk_level"] == "low"
    assert last["ok"] is True
    assert isinstance(last["duration_ms"], int)


def test_log_common_fields_for_excel_write(env):
    from agent import app

    src = env["work"] / "L2.xlsx"
    _make_xlsx(src, {"S": [["1"]]})
    app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": str(env["out"] / "w.xlsx"),
            "rows": [["a"], [1]],
        },
    )
    entries = _read_log(env["log"])
    last = entries[-1]
    assert _LOG_COMMON_KEYS.issubset(last.keys())
    assert last["action"] == "excel_write_report_copy"
    assert last["category"] == "excel"
    assert last["risk_level"] == "medium"
    assert last["ok"] is True


def test_log_does_not_leak_sensitive_words_for_excel(env):
    from agent import app

    src = env["work"] / "s.xlsx"
    _make_xlsx(src, {"S": [["token_value_here"]]})
    app.run("excel_read_sheet", options={"file_path": str(src)})
    app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": str(env["out"] / "ok.xlsx"),
            # 행 데이터는 로그에 기록되지 않는다
            "rows": [["password_leak_attempt"]],
        },
    )
    raw = env["log"].read_text(encoding="utf-8").lower()
    for banned in ("password", "token", "cookie", "authorization", "session"):
        assert banned not in raw, f"banned '{banned}' leaked in log"


# ══════════════════════════════════════════════════════════════════════
# 2단계: describe_workbook
# ══════════════════════════════════════════════════════════════════════
def test_describe_workbook_basic_success(env):
    from agent import app

    src = env["work"] / "wb1.xlsx"
    _make_xlsx(src, {"Sheet1": [["일자", "품명", "수량"], ["2026-04-01", "A", 2]]})

    out = app.run("excel_describe_workbook", options={"file_path": str(src)})
    assert set(out.keys()) == _COMMON_KEYS
    assert out["ok"] is True
    assert out["error"] is None
    assert out["action"] == "excel_describe_workbook"
    d = out["data"]
    assert d["sheet_count"] == 1
    assert len(d["sheets"]) == 1
    s0 = d["sheets"][0]
    assert s0["name"] == "Sheet1"
    assert s0["max_row"] >= 2
    assert s0["max_column"] >= 3
    assert s0["header_preview"][:3] == ["일자", "품명", "수량"]


def test_describe_workbook_multiple_sheets(env):
    from agent import app

    src = env["work"] / "wb2.xlsx"
    _make_xlsx(src, {
        "Alpha": [["a"]],
        "Beta": [["x", "y"], [1, 2]],
        "Gamma": [["헤더1", "헤더2", "헤더3"], [None, None, None]],
    })
    out = app.run("excel_describe_workbook", options={"file_path": str(src)})
    assert out["ok"] is True
    d = out["data"]
    assert d["sheet_count"] == 3
    names = [s["name"] for s in d["sheets"]]
    assert names == ["Alpha", "Beta", "Gamma"]
    # 각 시트에 max_row/max_column/header_preview 필드 존재
    for s in d["sheets"]:
        assert set(s.keys()) == {"name", "max_row", "max_column", "header_preview"}
        assert isinstance(s["header_preview"], list)


def test_describe_workbook_header_preview_picks_densest(env):
    """상단 5행 중 값이 가장 많이 채워진 행을 preview 로 사용."""
    from agent import app

    src = env["work"] / "wb3.xlsx"
    # 첫 행에는 빈 셀 포함, 두 번째 행에 더 많은 값이 존재.
    _make_xlsx(src, {"S": [
        ["주석", None, None],
        ["일자", "품명", "수량", "단가"],
        ["2026-04-01", "A", 1, 100],
    ]})
    out = app.run("excel_describe_workbook", options={"file_path": str(src)})
    assert out["ok"] is True
    preview = out["data"]["sheets"][0]["header_preview"]
    assert preview == ["일자", "품명", "수량", "단가"]


def test_describe_workbook_blocked_outside_work_dir(env):
    from agent import app, errors

    outside = env["tmp"] / "outside_wb.xlsx"
    _make_xlsx(outside, {"S": [["x"]]})
    out = app.run("excel_describe_workbook", options={"file_path": str(outside)})
    assert out["ok"] is False
    assert out["error"] == errors.FILE_NOT_ALLOWED


def test_describe_workbook_log_common_fields(env):
    from agent import app

    src = env["work"] / "wb4.xlsx"
    _make_xlsx(src, {"S": [["h"], [1]]})
    app.run("excel_describe_workbook", options={"file_path": str(src)})

    entries = _read_log(env["log"])
    last = entries[-1]
    assert _LOG_COMMON_KEYS.issubset(last.keys())
    assert last["action"] == "excel_describe_workbook"
    assert last["category"] == "excel"
    assert last["risk_level"] == "low"
    assert last["ok"] is True


# ══════════════════════════════════════════════════════════════════════
# 2단계: read_table
# ══════════════════════════════════════════════════════════════════════
def test_read_table_basic_success(env):
    from agent import app

    src = env["work"] / "rt1.xlsx"
    _make_xlsx(src, {"Sheet1": [
        ["일자", "품명", "수량", "단가"],
        ["2026-04-01", "A", 2, 1000],
        ["2026-04-02", "B", 3, 500],
    ]})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert set(out.keys()) == _COMMON_KEYS
    assert out["ok"] is True
    assert out["error"] is None
    assert out["action"] == "excel_read_table"
    d = out["data"]
    assert d["sheet_name"] == "Sheet1"
    assert d["header_row"] == 1
    assert d["columns"] == ["일자", "품명", "수량", "단가"]
    assert d["row_count"] == 2
    assert d["truncated"] is False
    assert d["rows"][0]["품명"] == "A"
    assert d["rows"][0]["수량"] == 2
    assert d["rows"][1]["단가"] == 500


def test_read_table_specific_sheet(env):
    from agent import app

    src = env["work"] / "rt2.xlsx"
    _make_xlsx(src, {
        "First": [["z"], [0]],
        "Target": [["a", "b"], [1, 2], [3, 4]],
    })
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "sheet_name": "Target"},
    )
    assert out["ok"] is True
    assert out["data"]["sheet_name"] == "Target"
    assert out["data"]["columns"] == ["a", "b"]
    assert out["data"]["row_count"] == 2


def test_read_table_custom_header_row(env):
    from agent import app

    src = env["work"] / "rt3.xlsx"
    _make_xlsx(src, {"S": [
        ["제목 줄 - 무시"],
        ["일자", "품명"],
        ["2026-04-01", "A"],
    ]})
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "header_row": 2},
    )
    assert out["ok"] is True
    assert out["data"]["header_row"] == 2
    assert out["data"]["columns"] == ["일자", "품명"]
    assert out["data"]["rows"] == [{"일자": "2026-04-01", "품명": "A"}]


def test_read_table_duplicate_headers_deduplicated(env):
    from agent import app

    src = env["work"] / "rt4.xlsx"
    _make_xlsx(src, {"S": [
        ["금액", "금액", "수량"],
        [100, 200, 5],
    ]})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert out["ok"] is True
    assert out["data"]["columns"] == ["금액", "금액_2", "수량"]
    assert out["data"]["rows"][0] == {"금액": 100, "금액_2": 200, "수량": 5}


def test_read_table_trailing_empty_columns_trimmed(env):
    from agent import app

    src = env["work"] / "rt5.xlsx"
    _make_xlsx(src, {"S": [
        ["a", "b", None, None],
        [1, 2, None, None],
    ]})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert out["ok"] is True
    assert out["data"]["columns"] == ["a", "b"]
    assert out["data"]["rows"][0] == {"a": 1, "b": 2}


def test_read_table_empty_header_row_blocked(env):
    from agent import app, errors

    src = env["work"] / "rt6.xlsx"
    _make_xlsx(src, {"S": [[None, None, None], [1, 2, 3]]})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert out["ok"] is False
    assert out["error"] == errors.TABLE_HEADER_NOT_FOUND


def test_read_table_numeric_only_header_blocked(env):
    from agent import app, errors

    # 숫자 헤더는 정책상 거부 (TABLE_HEADER_NOT_FOUND)
    src = env["work"] / "rt6b.xlsx"
    _make_xlsx(src, {"S": [[2024, 2025, 2026], [1, 2, 3]]})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert out["ok"] is False
    assert out["error"] == errors.TABLE_HEADER_NOT_FOUND


def test_read_table_empty_data_blocked(env):
    from agent import app, errors

    src = env["work"] / "rt7.xlsx"
    _make_xlsx(src, {"S": [["a", "b"]]})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert out["ok"] is False
    assert out["error"] == errors.TABLE_EMPTY


def test_read_table_default_max_rows_applied(env):
    from agent import app
    from agent.connectors import excel_connector as ec

    src = env["work"] / "rt8.xlsx"
    rows = [["a"]] + [[i] for i in range(50)]
    _make_xlsx(src, {"S": rows})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert out["ok"] is True
    assert out["data"]["row_count"] == 50
    # default 1000 상한이 적용됨을 확인 (상수 자체 검증)
    assert ec.DEFAULT_MAX_ROWS == 1000


def test_read_table_custom_max_rows(env):
    from agent import app

    src = env["work"] / "rt9.xlsx"
    rows = [["a"]] + [[i] for i in range(10)]
    _make_xlsx(src, {"S": rows})
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "max_rows": 3},
    )
    assert out["ok"] is True
    assert out["data"]["row_count"] == 3
    assert out["data"]["truncated"] is True


def test_read_table_truncated_marks_true_when_overflow(env):
    from agent import app

    src = env["work"] / "rt10.xlsx"
    rows = [["a"]] + [[i] for i in range(6)]
    _make_xlsx(src, {"S": rows})
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "max_rows": 5},
    )
    assert out["ok"] is True
    assert out["data"]["row_count"] == 5
    assert out["data"]["truncated"] is True


def test_read_table_no_truncation_when_exact(env):
    from agent import app

    src = env["work"] / "rt11.xlsx"
    rows = [["a"]] + [[i] for i in range(5)]
    _make_xlsx(src, {"S": rows})
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "max_rows": 5},
    )
    assert out["ok"] is True
    assert out["data"]["row_count"] == 5
    assert out["data"]["truncated"] is False


def test_read_table_invalid_max_rows_over_cap(env):
    from agent import app, errors

    src = env["work"] / "rt12.xlsx"
    _make_xlsx(src, {"S": [["a"], [1]]})
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "max_rows": 100_000},
    )
    assert out["ok"] is False
    assert out["error"] == errors.INVALID_MAX_ROWS


def test_read_table_invalid_max_rows_zero(env):
    from agent import app, errors

    src = env["work"] / "rt13.xlsx"
    _make_xlsx(src, {"S": [["a"], [1]]})
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "max_rows": 0},
    )
    assert out["ok"] is False
    assert out["error"] == errors.INVALID_MAX_ROWS


def test_read_table_invalid_header_row(env):
    from agent import app, errors

    src = env["work"] / "rt14.xlsx"
    _make_xlsx(src, {"S": [["a"], [1]]})
    for bad in (0, -1, "1"):
        out = app.run(
            "excel_read_table",
            options={"file_path": str(src), "header_row": bad},
        )
        assert out["ok"] is False, bad
        assert out["error"] == errors.INVALID_HEADER_ROW, bad


def test_read_table_missing_sheet_uses_shared_error(env):
    from agent import app, errors

    src = env["work"] / "rt15.xlsx"
    _make_xlsx(src, {"S": [["a"], [1]]})
    out = app.run(
        "excel_read_table",
        options={"file_path": str(src), "sheet_name": "Nope"},
    )
    assert out["ok"] is False
    assert out["error"] == errors.SHEET_NOT_FOUND


def test_read_table_outside_work_dir_blocked(env):
    from agent import app, errors

    outside = env["tmp"] / "ext.xlsx"
    _make_xlsx(outside, {"S": [["a"], [1]]})
    out = app.run("excel_read_table", options={"file_path": str(outside)})
    assert out["ok"] is False
    assert out["error"] == errors.FILE_NOT_ALLOWED


def test_read_table_skips_fully_empty_rows(env):
    from agent import app

    src = env["work"] / "rt16.xlsx"
    _make_xlsx(src, {"S": [
        ["a", "b"],
        [1, 2],
        [None, None],
        [3, 4],
    ]})
    out = app.run("excel_read_table", options={"file_path": str(src)})
    assert out["ok"] is True
    assert out["data"]["row_count"] == 2
    assert out["data"]["rows"] == [{"a": 1, "b": 2}, {"a": 3, "b": 4}]


def test_read_table_log_common_fields(env):
    from agent import app

    src = env["work"] / "rt17.xlsx"
    _make_xlsx(src, {"S": [["a"], [1], [2]]})
    app.run("excel_read_table", options={"file_path": str(src)})
    entries = _read_log(env["log"])
    last = entries[-1]
    assert _LOG_COMMON_KEYS.issubset(last.keys())
    assert last["action"] == "excel_read_table"
    assert last["category"] == "excel"
    assert last["risk_level"] == "low"
    assert last["ok"] is True
    assert last.get("truncated") is False


# ══════════════════════════════════════════════════════════════════════
# 2단계: 신규 에러 코드 등록 확인
# ══════════════════════════════════════════════════════════════════════
def test_phase2_error_codes_registered():
    from agent import errors
    for name in (
        "TABLE_HEADER_NOT_FOUND",
        "TABLE_EMPTY",
        "TABLE_TOO_LARGE",
        "INVALID_HEADER_ROW",
        "INVALID_MAX_ROWS",
    ):
        assert hasattr(errors, name), f"errors.{name} missing"
        assert getattr(errors, name) in errors.ALL_CODES, name
        assert errors.is_standard_code(getattr(errors, name))


# ══════════════════════════════════════════════════════════════════════
# 2단계: 1단계 회귀 방지 — read_sheet / write_report_copy 그대로 동작
# ══════════════════════════════════════════════════════════════════════
def test_phase1_read_sheet_regression(env):
    from agent import app

    src = env["work"] / "reg1.xlsx"
    _make_xlsx(src, {"Sheet1": [["h"], [1]]})
    out = app.run("excel_read_sheet", options={"file_path": str(src)})
    assert out["ok"] is True
    assert out["data"]["row_count"] == 2
    assert out["data"]["sheet_name"] == "Sheet1"


def test_phase1_write_report_copy_regression(env):
    from agent import app

    src = env["work"] / "reg2.xlsx"
    _make_xlsx(src, {"S": [["x"]]})
    target = env["out"] / "reg_out.xlsx"
    out = app.run(
        "excel_write_report_copy",
        options={
            "file_path": str(src),
            "output_path": str(target),
            "rows": [["h"], [1]],
        },
    )
    assert out["ok"] is True
    assert target.exists()


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
