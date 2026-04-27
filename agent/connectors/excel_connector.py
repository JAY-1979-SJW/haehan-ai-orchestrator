"""Excel 커넥터 (1단계).

지원 action:
- ``excel_read_sheet(file_path, sheet_name=None, range_ref=None)``
- ``excel_write_report_copy(file_path, output_path, rows)``

원칙:
- 원본 파일 overwrite 금지 — 쓰기 action 은 반드시 별도 경로로만 저장.
- 허용 경로 검증은 호출자(app/runner) 가 ``agent.file_policy`` 로 선행 수행.
  이 모듈은 이미 검증된 ``pathlib.Path`` 를 받는다.
- openpyxl 기반. ``.xlsm`` 등 제약 포맷은 명시적 에러 코드 반환.
- 회원가입/제출/네트워크 호출 없음. 로컬 파일 I/O 에 한정.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

from .. import errors as _err

_SUPPORTED_READ_EXTS = frozenset({".xlsx"})
_SUPPORTED_WRITE_EXTS = frozenset({".xlsx"})

# ── 2단계 가드 상수 ────────────────────────────────────────────────────
# describe_workbook: 각 시트 상단 몇 행까지만 미리보기에 사용할지.
HEADER_SCAN_ROWS = 5
# header_preview 로 반환할 최대 셀 수 (앞에서부터).
HEADER_PREVIEW_MAX_CELLS = 20

# read_table 상한. 요청 max_rows 는 [1, MAX_ROWS_CAP] 범위.
DEFAULT_MAX_ROWS = 1000
MAX_ROWS_CAP = 5000
# 시트 폭이 이보다 크면 TABLE_TOO_LARGE 로 거부.
MAX_COLUMNS_GUARD = 1000
# 데이터 스캔 시 max_rows 초과 구간에서 허용되는 조회 여유 (빈 행 포함).
SCAN_OVERSHOOT = 2000


def _check_ext(path: Path, allowed: frozenset[str]) -> Optional[str]:
    ext = path.suffix.lower()
    if ext not in allowed:
        return _err.EXCEL_UNSUPPORTED_FORMAT
    return None


def _parse_range(range_ref: str) -> Optional[tuple[int, int, int, int]]:
    """``A1:B3`` 형태의 범위를 ``(min_row, min_col, max_row, max_col)`` 로.

    지원되지 않으면 None. openpyxl 의 ``range_boundaries`` 를 쓰되
    실패 시 조용히 None 을 반환해 상위에서 에러 처리.
    """
    if not isinstance(range_ref, str) or not range_ref.strip():
        return None
    try:
        from openpyxl.utils.cell import range_boundaries  # type: ignore
    except ImportError:
        return None
    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_ref.strip())
    except (ValueError, TypeError):
        return None
    if None in (min_col, min_row, max_col, max_row):
        return None
    return int(min_row), int(min_col), int(max_row), int(max_col)


def excel_read_sheet(
    file_path: Path,
    sheet_name: Optional[str] = None,
    range_ref: Optional[str] = None,
) -> tuple[Optional[dict], Optional[str]]:
    """지정 시트를 읽어 rows (2D list) 로 반환.

    반환: (data_or_None, error_or_None).
    data = {
        "file_path": str,
        "sheet_name": str,
        "rows": list[list[Any]],
        "row_count": int,
        "column_count": int,
    }
    """
    fmt_err = _check_ext(file_path, _SUPPORTED_READ_EXTS)
    if fmt_err:
        return None, fmt_err

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return None, _err.with_reason(_err.EXCEL_ERROR, "openpyxl_not_installed")

    try:
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return None, _err.with_reason(_err.EXCEL_ERROR, type(e).__name__)

    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                return None, _err.SHEET_NOT_FOUND
            ws = wb[sheet_name]
        else:
            ws = wb.active

        resolved_sheet = ws.title

        if range_ref:
            bounds = _parse_range(range_ref)
            if bounds is None:
                return None, _err.with_reason(_err.EXCEL_ERROR, "bad_range_ref")
            min_row, min_col, max_row, max_col = bounds
            rows_iter = ws.iter_rows(
                min_row=min_row,
                min_col=min_col,
                max_row=max_row,
                max_col=max_col,
                values_only=True,
            )
        else:
            rows_iter = ws.iter_rows(values_only=True)

        rows: list[list[Any]] = [list(r) for r in rows_iter]
    except Exception as e:  # noqa: BLE001
        return None, _err.with_reason(_err.EXCEL_ERROR, type(e).__name__)
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass

    row_count = len(rows)
    column_count = max((len(r) for r in rows), default=0)

    data = {
        "file_path": str(file_path),
        "sheet_name": resolved_sheet,
        "rows": rows,
        "row_count": row_count,
        "column_count": column_count,
    }
    return data, None


def excel_write_report_copy(
    source_file_path: Path,
    output_path: Path,
    rows: list[list[Any]],
) -> tuple[Optional[dict], Optional[str]]:
    """``rows`` 를 새 workbook 에 기록한 뒤 ``output_path`` 에 저장.

    원본 파일은 읽지도/변경하지도 않는다. 이 함수는 source_file_path 를
    "참고 식별자" 로만 사용하며 실제 파일 조작 대상은 output_path 뿐이다.

    호출 전 검증 전제 (호출자 책임):
    - source_file_path 는 AGENT_WORK_DIR 하위에 존재하는 파일
    - output_path 는 AGENT_OUTPUT_DIR 하위 + 아직 존재하지 않음
    """
    if not isinstance(rows, list):
        return None, _err.ROWS_REQUIRED

    fmt_err = _check_ext(output_path, _SUPPORTED_WRITE_EXTS)
    if fmt_err:
        return None, fmt_err

    try:
        from openpyxl import Workbook  # type: ignore
    except ImportError:
        return None, _err.with_reason(_err.EXCEL_ERROR, "openpyxl_not_installed")

    written = 0
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "report"
        for row in rows:
            if not isinstance(row, (list, tuple)):
                ws.append([row])
            else:
                ws.append(list(row))
            written += 1
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
    except Exception as e:  # noqa: BLE001
        return None, _err.with_reason(_err.EXCEL_ERROR, type(e).__name__)

    data = {
        "source_file_path": str(source_file_path),
        "output_path": str(output_path),
        "written_row_count": written,
    }
    return data, None


def _is_empty_cell(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and not v.strip():
        return True
    return False


def _preview_cell(v: Any) -> Optional[str]:
    """preview 에 담을 값으로 정규화. 빈 값은 None 으로 표시해 호출자가 drop."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    try:
        return str(v)
    except Exception:  # noqa: BLE001
        return None


def _extract_header_preview(ws: Any) -> list[str]:
    """상단 ``HEADER_SCAN_ROWS`` 행 중 값이 가장 많이 채워진 한 행을 골라
    앞에서부터 최대 ``HEADER_PREVIEW_MAX_CELLS`` 개 셀을 반환.

    read_only iterator 로 앞부분만 스캔한다.
    """
    best: list[str] = []
    try:
        rows_iter = ws.iter_rows(
            min_row=1,
            max_row=HEADER_SCAN_ROWS,
            max_col=HEADER_PREVIEW_MAX_CELLS,
            values_only=True,
        )
    except TypeError:
        rows_iter = ws.iter_rows(values_only=True)
    for i, row in enumerate(rows_iter):
        if i >= HEADER_SCAN_ROWS:
            break
        cells = [_preview_cell(c) for c in row[:HEADER_PREVIEW_MAX_CELLS]]
        compact = [c for c in cells if c is not None]
        if len(compact) > len(best):
            best = compact
    return best


def describe_workbook(file_path: Path) -> tuple[Optional[dict], Optional[str]]:
    """Workbook 구조를 요약 반환.

    반환: (data_or_None, error_or_None).
    data = {
        "sheet_count": int,
        "sheets": [
            {"name": str, "max_row": int, "max_column": int,
             "header_preview": list[str]}
        ],
    }
    """
    fmt_err = _check_ext(file_path, _SUPPORTED_READ_EXTS)
    if fmt_err:
        return None, fmt_err

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return None, _err.with_reason(_err.EXCEL_ERROR, "openpyxl_not_installed")

    try:
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return None, _err.with_reason(_err.EXCEL_ERROR, type(e).__name__)

    sheets: list[dict] = []
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            max_row_raw = ws.max_row if ws.max_row is not None else 0
            max_col_raw = ws.max_column if ws.max_column is not None else 0
            try:
                max_row = int(max_row_raw)
                max_column = int(max_col_raw)
            except (TypeError, ValueError):
                max_row, max_column = 0, 0

            if max_column > MAX_COLUMNS_GUARD:
                return None, _err.TABLE_TOO_LARGE

            header_preview = _extract_header_preview(ws)
            sheets.append({
                "name": name,
                "max_row": max_row,
                "max_column": max_column,
                "header_preview": header_preview,
            })
    except Exception as e:  # noqa: BLE001
        return None, _err.with_reason(_err.EXCEL_ERROR, type(e).__name__)
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass

    return {
        "file_path": str(file_path),
        "sheet_count": len(sheets),
        "sheets": sheets,
    }, None


def _normalize_headers(raw: list[Any]) -> tuple[list[Optional[str]], list[str]]:
    """헤더 셀 리스트를 (slot_names, column_names) 로 분해.

    - slot_names[i] = i번째 셀의 컬럼명 (빈 헤더 자리면 None)
    - column_names = None 을 제거한 컬럼명의 순서 리스트 (중복 보정 포함)
    - trailing empty 셀은 호출자가 미리 제거한 뒤 전달.
    """
    seen: dict[str, int] = {}
    slot_names: list[Optional[str]] = []
    column_names: list[str] = []
    for cell in raw:
        if not isinstance(cell, str) or not cell.strip():
            slot_names.append(None)
            continue
        name = cell.strip()
        if name not in seen:
            seen[name] = 1
            slot_names.append(name)
            column_names.append(name)
        else:
            seen[name] += 1
            dup = f"{name}_{seen[name]}"
            slot_names.append(dup)
            column_names.append(dup)
    return slot_names, column_names


def read_table(
    file_path: Path,
    sheet_name: Optional[str] = None,
    header_row: Optional[int] = 1,
    max_rows: Optional[int] = None,
) -> tuple[Optional[dict], Optional[str]]:
    """헤더 기반으로 시트를 읽어 JSON rows 로 반환.

    반환: (data_or_None, error_or_None).
    data = {
        "sheet_name": str,
        "header_row": int,
        "columns": list[str],
        "row_count": int,
        "rows": list[dict],
        "truncated": bool,
    }
    """
    # 입력 정규화/검증
    if header_row is None:
        header_row = 1
    if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
        return None, _err.INVALID_HEADER_ROW

    if max_rows is None:
        max_rows = DEFAULT_MAX_ROWS
    if (
        isinstance(max_rows, bool)
        or not isinstance(max_rows, int)
        or max_rows < 1
        or max_rows > MAX_ROWS_CAP
    ):
        return None, _err.INVALID_MAX_ROWS

    fmt_err = _check_ext(file_path, _SUPPORTED_READ_EXTS)
    if fmt_err:
        return None, fmt_err

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return None, _err.with_reason(_err.EXCEL_ERROR, "openpyxl_not_installed")

    try:
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return None, _err.with_reason(_err.EXCEL_ERROR, type(e).__name__)

    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                return None, _err.SHEET_NOT_FOUND
            ws = wb[sheet_name]
        else:
            ws = wb.active
        resolved_sheet = ws.title

        # 시트 폭 가드 (가능한 경우)
        try:
            mc = int(ws.max_column) if ws.max_column is not None else 0
        except (TypeError, ValueError):
            mc = 0
        if mc > MAX_COLUMNS_GUARD:
            return None, _err.TABLE_TOO_LARGE

        # 헤더 읽기
        header_tuple = None
        for row in ws.iter_rows(
            min_row=header_row, max_row=header_row, values_only=True,
        ):
            header_tuple = row
            break
        if header_tuple is None:
            return None, _err.TABLE_HEADER_NOT_FOUND

        raw_headers = list(header_tuple)
        while raw_headers and _is_empty_cell(raw_headers[-1]):
            raw_headers.pop()
        if not raw_headers:
            return None, _err.TABLE_HEADER_NOT_FOUND
        if not any(isinstance(c, str) and c.strip() for c in raw_headers):
            return None, _err.TABLE_HEADER_NOT_FOUND

        slot_names, column_names = _normalize_headers(raw_headers)
        usable_width = len(slot_names)

        # 데이터 행 읽기 — 헤더 바로 아래부터.
        rows_out: list[dict] = []
        truncated = False
        data_start = header_row + 1
        # 하드 스캔 상한: max_rows + SCAN_OVERSHOOT 행까지만 조회 (빈 행 포함).
        hard_stop = data_start + max_rows + SCAN_OVERSHOOT - 1

        scanned = 0
        for row in ws.iter_rows(
            min_row=data_start,
            max_row=hard_stop,
            values_only=True,
        ):
            scanned += 1
            cells = list(row)[:usable_width] if usable_width else list(row)
            if all(_is_empty_cell(c) for c in cells) or not cells:
                continue
            if len(rows_out) >= max_rows:
                truncated = True
                break
            entry: dict = {}
            for idx, name in enumerate(slot_names):
                if name is None:
                    continue
                entry[name] = cells[idx] if idx < len(cells) else None
            rows_out.append(entry)
        del scanned  # 가드용 변수, 값 자체는 반환하지 않는다.

        if not rows_out and not truncated:
            return None, _err.TABLE_EMPTY
    except Exception as e:  # noqa: BLE001
        return None, _err.with_reason(_err.EXCEL_ERROR, type(e).__name__)
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass

    return {
        "file_path": str(file_path),
        "sheet_name": resolved_sheet,
        "header_row": header_row,
        "columns": column_names,
        "row_count": len(rows_out),
        "rows": rows_out,
        "truncated": truncated,
    }, None


__all__ = [
    "excel_read_sheet",
    "excel_write_report_copy",
    "describe_workbook",
    "read_table",
    "DEFAULT_MAX_ROWS",
    "MAX_ROWS_CAP",
    "MAX_COLUMNS_GUARD",
    "HEADER_PREVIEW_MAX_CELLS",
]
