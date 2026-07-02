from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from .models import AttachmentResult

TEXT_LIMIT = 250_000


def _clip(text: str, limit: int = TEXT_LIMIT) -> str:
    return re.sub(r"\s+", " ", text or "").strip()[:limit]


def parse_attachment(path: str | Path) -> AttachmentResult:
    file_path = Path(path)
    result = AttachmentResult(
        filename=file_path.name,
        path=str(file_path),
        size_bytes=file_path.stat().st_size if file_path.exists() else None,
    )
    try:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            result.parser = "pymupdf"
            result.text = parse_pdf(file_path)
        elif suffix == ".hwpx":
            result.parser = "hwpx-xml"
            result.text = parse_hwpx(file_path)
        elif suffix in {".xlsx", ".xlsm"}:
            result.parser = "openpyxl"
            result.text = parse_xlsx(file_path)
        elif suffix == ".zip":
            result.parser = "zip-recursive"
            result.text = parse_zip(file_path)
        elif suffix in {".txt", ".md", ".csv", ".html", ".htm"}:
            result.parser = "text"
            result.text = _clip(file_path.read_text(encoding="utf-8", errors="ignore"))
        elif suffix == ".hwp":
            result.parser = "hwp-unsupported"
            result.error = "HWP 바이너리는 HWPX 변환 후 재분석이 필요합니다."
        else:
            result.parser = "unsupported"
            result.error = f"지원하지 않는 첨부 형식입니다: {suffix}"
    except Exception as exc:
        result.error = f"{type(exc).__name__}: {exc}"
    return result


def parse_pdf(path: Path) -> str:
    try:
        import fitz  # type: ignore
    except Exception:
        return ""
    chunks: list[str] = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            chunks.append(page.get_text("text"))
    return _clip("\n".join(chunks))


def parse_hwpx(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as zf:
        for name in sorted(zf.namelist()):
            lower = name.lower()
            if not lower.endswith(".xml"):
                continue
            if "contents/" not in lower and "section" not in lower:
                continue
            try:
                root = ET.fromstring(zf.read(name))
            except ET.ParseError:
                continue
            for node in root.iter():
                if node.text and node.text.strip():
                    chunks.append(node.text.strip())
    return _clip("\n".join(chunks))


def parse_xlsx(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: list[str] = []
    try:
        for ws in wb.worksheets:
            rows.append(f"[시트] {ws.title}")
            for row in ws.iter_rows(values_only=True):
                values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if values:
                    rows.append(" | ".join(values))
                if len("\n".join(rows)) > TEXT_LIMIT:
                    break
    finally:
        wb.close()
    return _clip("\n".join(rows))


def parse_zip(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            suffix = Path(name).suffix.lower()
            if suffix not in {".txt", ".md", ".csv", ".hwpx"}:
                continue
            chunks.append(f"[압축 내부 파일] {name}")
            if suffix in {".txt", ".md", ".csv"}:
                chunks.append(zf.read(name).decode("utf-8", errors="ignore"))
            elif suffix == ".hwpx":
                chunks.append("HWPX 내부 파일은 압축 밖으로 풀어서 상세 분석 권장")
    return _clip("\n".join(chunks))
